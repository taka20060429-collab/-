# -*- coding: utf-8 -*-
"""pptx を読み、各スライドを PNG にレンダリングして確認用プレビューを作る簡易レンダラ。
python-pptx の図形座標・テキスト・色を使って Pillow で描画する（近似）。"""
import sys
from pptx import Presentation
from pptx.util import Emu
from PIL import Image, ImageDraw, ImageFont

FONT_PATH = "/usr/share/fonts/opentype/ipafont-gothic/ipagp.ttf"
SCALE = 110  # px per inch
EMU_PER_INCH = 914400

prs = Presentation(sys.argv[1] if len(sys.argv) > 1 else "/tmp/deck.pptx")
W = int(prs.slide_width / EMU_PER_INCH * SCALE)
H = int(prs.slide_height / EMU_PER_INCH * SCALE)

_font_cache = {}
def font(pt, bold=False):
    key = (int(pt), bold)
    if key not in _font_cache:
        px = max(8, int(pt * SCALE / 72))
        _font_cache[key] = ImageFont.truetype(FONT_PATH, px)
    return _font_cache[key]

def emu_px(v):
    return int(v / EMU_PER_INCH * SCALE)

def rgb(c):
    try:
        return (c[0], c[1], c[2])
    except Exception:
        return (0, 0, 0)

def wrap(draw, text, fnt, max_w):
    if not text:
        return [""]
    lines, cur = [], ""
    for ch in text:
        if ch == "\n":
            lines.append(cur); cur = ""; continue
        t = cur + ch
        if draw.textlength(t, font=fnt) > max_w and cur:
            lines.append(cur); cur = ch
        else:
            cur = t
    lines.append(cur)
    return lines

def render(slide, idx):
    img = Image.new("RGB", (W, H), (255, 255, 255))
    d = ImageDraw.Draw(img)
    for shp in slide.shapes:
        try:
            x, y = emu_px(shp.left), emu_px(shp.top)
            w, h = emu_px(shp.width), emu_px(shp.height)
        except Exception:
            continue
        # 図形塗り
        if shp.shape_type is not None and not shp.has_text_frame or (hasattr(shp, "fill")):
            pass
        try:
            if shp.has_table:
                tbl = shp.table
                nrows = len(tbl.rows); ncols = len(tbl.columns)
                col_w = [emu_px(c.width) for c in tbl.columns]
                row_h = h // max(nrows, 1)
                cy = y
                for r in range(nrows):
                    cx = x
                    for c in range(ncols):
                        cell = tbl.cell(r, c)
                        fc = None
                        try:
                            fc = rgb(cell.fill.fore_color.rgb)
                        except Exception:
                            fc = (255, 255, 255)
                        d.rectangle([cx, cy, cx + col_w[c], cy + row_h], fill=fc, outline=(200, 200, 200))
                        para = cell.text_frame.paragraphs[0]
                        sz = para.font.size.pt if para.font.size else 14
                        bd = bool(para.font.bold)
                        try:
                            tcol = rgb(para.font.color.rgb)
                        except Exception:
                            tcol = (0, 0, 0)
                        fnt = font(sz, bd)
                        d.text((cx + 6, cy + 6), cell.text, fill=tcol, font=fnt)
                        cx += col_w[c]
                    cy += row_h
                continue
        except Exception:
            pass
        # 単純図形（塗り）
        try:
            if shp.shape_type == 1:  # rectangle (autoshape)
                fc = rgb(shp.fill.fore_color.rgb)
                d.rectangle([x, y, x + w, y + h], fill=fc)
        except Exception:
            pass
        # テキスト
        if shp.has_text_frame:
            ty = y + 4
            for para in shp.text_frame.paragraphs:
                txt = "".join(run.text for run in para.runs) or para.text
                if not para.runs and not txt:
                    ty += int(12 * SCALE / 72); continue
                run = para.runs[0] if para.runs else None
                sz = (run.font.size.pt if run and run.font.size else 18)
                bd = bool(run.font.bold) if run else False
                try:
                    col = rgb(run.font.color.rgb) if run and run.font.color and run.font.color.type is not None else (34, 34, 34)
                except Exception:
                    col = (34, 34, 34)
                fnt = font(sz, bd)
                align = para.alignment
                for line in wrap(d, txt, fnt, w - 8):
                    tw = d.textlength(line, font=fnt)
                    lx = x + 4
                    if str(align) == "CENTER (2)":
                        lx = x + (w - tw) / 2
                    elif str(align) == "RIGHT (3)":
                        lx = x + w - tw - 4
                    d.text((lx, ty), line, fill=col, font=fnt)
                    ty += int(sz * SCALE / 72 * 1.25)
    img.save(f"/tmp/preview/slide_{idx:02d}.png")

import os
os.makedirs("/tmp/preview", exist_ok=True)
for i, sl in enumerate(prs.slides, 1):
    render(sl, i)
print("rendered", len(prs.slides._sldIdLst), "slides")
