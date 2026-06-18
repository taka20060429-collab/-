# -*- coding: utf-8 -*-
"""ネットワーク実習 実験レポート（文書/Word）生成スクリプト
過去レポート(ネットワーク実習・出口和佳奈)の構成を踏襲：目的→考察(課題1〜7)→まとめ→参考文献。
黒一色・A4縦。結果は記入シート写真を表紙の次に貼付。
"""
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

IMGDIR = os.path.join(os.path.dirname(__file__), "img")
BODY_FONT = "ＭＳ 明朝"
HEAD_FONT = "ＭＳ ゴシック"
ASCII_FONT = "Times New Roman"

doc = Document()
style = doc.styles["Normal"]
style.font.name = ASCII_FONT
style.font.size = Pt(10.5)
style.font.color.rgb = RGBColor(0, 0, 0)
style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
style.paragraph_format.line_spacing = 1.15
style.paragraph_format.space_after = Pt(0)

sec = doc.sections[0]
sec.page_width = Mm(210); sec.page_height = Mm(297)
sec.top_margin = Mm(25); sec.bottom_margin = Mm(25)
sec.left_margin = Mm(25); sec.right_margin = Mm(25)


def set_run_font(run, ascii_font=ASCII_FONT, ea_font=BODY_FONT, size=None, bold=False):
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    rf.set(qn("w:ascii"), ascii_font); rf.set(qn("w:hAnsi"), ascii_font)
    rf.set(qn("w:eastAsia"), ea_font)
    if size:
        run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def para(text="", size=10.5, bold=False, align=None, ea=BODY_FONT, ascii_f=ASCII_FONT,
         space_before=0, space_after=4, indent=None, line=1.3):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before); pf.space_after = Pt(space_after); pf.line_spacing = line
    if indent is not None:
        pf.first_line_indent = Pt(indent)
    if text:
        r = p.add_run(text); set_run_font(r, ascii_f, ea, size, bold)
    return p


def heading(num, text, level=1):
    sizes = {1: 14, 2: 12}; sb = {1: 14, 2: 8}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(sb[level]); p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    label = f"{num}　{text}" if text else num
    r = p.add_run(label); set_run_font(r, HEAD_FONT, HEAD_FONT, sizes[level], bold=True)
    return p


def body(text, indent=10.5):
    return para(text, size=10.5, space_after=5, indent=indent, line=1.35)


def item(text):
    """箇条書き（・）。字下げ無しの行頭記号。"""
    p = para("・" + text, size=10.5, space_after=3, indent=None, line=1.3)
    p.paragraph_format.left_indent = Mm(4)
    return p


def add_figure(path, caption, width_mm=160):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    if os.path.exists(path):
        run.add_picture(path, width=Mm(width_mm))
    else:
        rr = p.add_run(f"[画像なし: {path}]"); set_run_font(rr)
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    cr = cap.add_run(caption); set_run_font(cr, ASCII_FONT, BODY_FONT, 9.5)


def page_number_footer():
    p = sec.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    ins = OxmlElement("w:instrText"); ins.set(qn("xml:space"), "preserve"); ins.text = "PAGE"
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    run._r.append(f1); run._r.append(ins); run._r.append(f2)
    set_run_font(run, ASCII_FONT, BODY_FONT, 10)


# =================================================================== 表紙
for _ in range(3):
    para("", space_after=2)
para("電気基礎実験", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
     ea=HEAD_FONT, ascii_f=HEAD_FONT, space_after=10)
para("ネットワーク実習", size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
     ea=HEAD_FONT, ascii_f=HEAD_FONT, space_after=60)
para("実験者氏名：4326091　森木崇允", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
para("共同実験者：4326033　諸岡翔大", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
para("実験日：【要記入】", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
para("提出日：【要記入】", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
doc.add_page_break()

# =================================================================== 結果シート（表紙の次）
para("実験結果（レポート提出課題 記入シート）を以下に示す。", size=10.5, space_after=4)
add_figure(os.path.join(IMGDIR, "fig_results.jpg"),
           "図1　実験結果（レポート提出課題の記入シート）", width_mm=160)
doc.add_page_break()

# =================================================================== 1 実験目的
heading("1", "実験目的")
body("現代の人間社会において、情報通信の基盤となっている不可欠なシステムがインターネットである。"
     "インターネットはその機能性や利便性の反面、目に見えない危険が常に潜んでいる。そのため、利用者"
     "はその仕組みやリスクを正しく理解したうえで、自分で自分の身を守る必要がある。")
body("本実習では、ネットワークケーブルのうち最も一般的に使用されているシールド無ツイストペア"
     "（UTP：Unshielded Twisted Pair）ケーブルについて学び、実際にケーブル製作と接続性テストを行う"
     "ことを目的とする。また、作製したUTPケーブルとブロードバンドルータを用いて有線／無線両方の"
     "ネットワークを構築し、セキュリティ選定を行う。")

# =================================================================== 2 考察
heading("2", "考察")

heading("2-1", "考察課題1", level=2)
body("インターネットとはどのようなものかを説明する。")
body("(1) インターネットとは　インターネット（Internet）とは、言葉のとおりコンピュータ等の情報機器"
     "の「間（＝inter）」に張り巡らされた「網目（＝net）」のことである。より具体的には、世界中の"
     "情報機器を接続するネットワーク、すなわちデータ通信網のことである。インターネットは1960年代に"
     "アメリカで開発されたパケット通信ネットワークであるARPANET（Advanced Research Projects Agency "
     "Network）を起源とし、1990年代ごろから世界中で広く使われ始めた。その後、瞬く間に成長し、現在では"
     "人間社会を運用する基盤として必要不可欠な存在となっている。")
body("(2) インターネットの仕組み　各コンピュータは、ルータと呼ばれる機器や通信事業者の設置した機器"
     "を介し、有線もしくは無線でつながれている。同一のルータ等に接続されている機器が集まる家庭・"
     "会社・学校などの1つのネットワーク同士をつなげるのもインターネットの仕組みである。情報のやり取り"
     "には、共通のルールであるTCP／IP（Transmission Control Protocol／Internet Protocol）という"
     "プロトコルが用いられる。各機器はIPアドレスによって識別され、ルータが経路制御（ルーティング）を"
     "行うことで通信が成立する。インターネット上に設置されたWebサーバやメールサーバといった多種多様な"
     "サーバが、このプロトコルに則ってクライアントからの指示に応じた情報の送受信を行っている。")
body("(3) インターネットの危険性　インターネットは非常に便利で革新的なシステムである一方、不正"
     "アクセス、ウイルス感染、個人情報の漏洩、なりすまし、フィッシング詐欺など、目に見えない多くの"
     "脅威が常に潜んでいる。利用者はこれらのリスクを正しく理解し、通信の暗号化やパスワードの適切な"
     "管理、ソフトウェアの更新などの対策を講じて、自分の身を守る必要がある。")

heading("2-2", "考察課題2", level=2)
body("物理アドレス（MACアドレス）と論理アドレス（IPアドレス）とは、それぞれどのようなものかを説明"
     "する。")
body("物理アドレス（MACアドレス）は、ネットワークインターフェースカード（NIC）に製造時に割り当て"
     "られる48ビットの固有のアドレスであり、通常は16進数で表記される。データリンク層において、同一"
     "ネットワーク内の機器を一意に識別するために用いられ、原則として変更されないハードウェア固有の"
     "値である。本実習でも、各PCの有線・無線NICに固有のMACアドレスが割り当てられていることを確認した。")
body("一方、論理アドレス（IPアドレス）は、ネットワーク層で用いられるアドレスであり、ネットワーク"
     "管理者やDHCPサーバによって割り当てられ、ネットワークの構成に応じて変更することができる。IPv4"
     "では32ビットで表され、ネットワーク部とホスト部から構成される。MACアドレスが物理的・固定的で"
     "同一ネットワーク内の識別に用いられるのに対し、IPアドレスは論理的・可変であり、異なるネットワーク"
     "間の通信（ルーティング）に用いられる点が両者の大きな違いである。")

heading("2-3", "考察課題3", level=2)
body("ネットマスクとデフォルトゲートウェイとはどのようなものかを説明する。")
body("ネットマスク（サブネットマスク）は、IPアドレスのうちどこまでがネットワーク部で、どこからが"
     "ホスト部かを示す32ビットの値であり、一般的には255.255.255.0のように10進数で表記される。これに"
     "よって、通信相手が自分と同一のネットワーク（サブネット）に属しているか否かを判定することができる。")
body("デフォルトゲートウェイは、自分の属するネットワークの外部（別のネットワークやインターネット）"
     "へ通信する際に、パケットを最初に渡す中継機器（通常はルータ）のIPアドレスである。宛先が同一"
     "サブネット内であれば直接通信し、サブネット外であればデフォルトゲートウェイへパケットを送る"
     "ことで、外部ネットワークとの通信が可能になる。本実習でも、有線・無線接続時にデフォルトゲート"
     "ウェイとして192.168.11.1が割り当てられていることを確認した。")

heading("2-4", "考察課題4", level=2)
body("ネットワークケーブルの種類と特徴を述べる。")
body("・ツイストペアケーブル　絶縁被覆された銅線を2本ずつ撚り合わせて対にしたケーブルで、撚りに"
     "よって電磁干渉（EMI）を打ち消す効果がある。撚りの密度が高いほど高速で干渉に強く、CAT3、CAT5、"
     "CAT5e、CAT6などの規格がある。シールドのないUTPは安価で配線も容易なため最も広く用いられる。"
     "ノイズの多い環境では、シールド付きのSTP（Shielded Twisted-Pair）やScTP（Screened Twisted-Pair）"
     "が用いられるが、高価で柔軟性が低く扱いにくい。本実習ではUTPケーブルを用いてストレートケーブル"
     "とクロスケーブルを製作した。")
body("・同軸ケーブル　中心導体を絶縁体と網状の外部導体で覆った構造のケーブルで、古くはLANやテレビ"
     "放送などに用いられた。")
body("・光ファイバケーブル　光信号によって情報を伝送するため、超広帯域・長距離伝送が可能であり、"
     "電磁干渉の影響を受けず、傍受もされにくい。ただし価格が高く、曲げに弱く加工が難しい。")
body("なお、ツイストペアケーブルの結線にはTIA／EIAのT568AとT568Bの2つの規格があり、両端を同じ規格"
     "にしたものをストレートケーブル、一方の端をT568A・他方をT568Bにしたものをクロスケーブルと呼ぶ。")

heading("2-5", "考察課題5", level=2)
body("有線LANの利点と欠点を述べる。")
body("利点として、通信が安定して高速であり遅延が小さいこと、電波干渉を受けないこと、ケーブルで"
     "物理的に接続するため外部からの傍受や侵入がされにくくセキュリティ面で有利であることが挙げられる。")
body("欠点として、ケーブルの敷設が必要なため機器の移動や増設がしにくいこと、配線のスペースやコスト"
     "がかかること、ケーブル長に制限があること（イーサネットのUTPケーブルは規格上最大100m程度）が"
     "挙げられる。")

heading("2-6", "考察課題6", level=2)
body("無線LANの利点と欠点を述べる。")
body("利点として、ケーブルが不要なため配線の手間がなく、対応機器を電波の届く範囲で自由に移動・増設"
     "でき、見た目もすっきりすることが挙げられる。")
body("欠点として、電波干渉や障害物・距離によって通信速度や安定性が低下しやすく、一般に有線より低速"
     "で遅延も大きいこと、電波が傍受されうるためセキュリティ対策が必須であること、同時接続数や帯域"
     "にも制約があることが挙げられる。")

heading("2-7", "考察課題7", level=2)
body("無線LANのセキュリティを向上させる方法を示す。")
body("考察課題6で述べたように、無線LANはオフィスや学校等において、ケーブルの配線がいらないことに"
     "よる様々なメリットをもたらす。しかしその反面、機器から電波を飛ばして端末をつなぐ仕組みによる"
     "リスクも生じる。そこで、その恩恵を最大限に受けるためには、不正傍受や通信妨害、データの改ざんや"
     "破壊といった脅威を防ぐセキュリティ対策をしっかり行ったうえで、無線LANを活用する必要がある。"
     "以下に、企業などで行うべきセキュリティ対策方法を挙げる。")
item("WPA2／WPA3などの強力な方式を用いて通信を暗号化する（脆弱なWEPは避ける）")
item("暗号化キー（パスフレーズ）は推測されにくい強固なものを使用し、定期的に変更する")
item("来客用など外部の人が使う無線LAN環境は、本来のネットワークと分離して構築する")
item("MACアドレス認証（フィルタリング）を利用し、接続できる機器を制限する")
item("電波遮断シートの利用やアクセスポイントの設置場所の工夫など、物理的な対策を講じる")
item("ルータの管理者パスワードを初期値から変更し、ファームウェアを最新の状態に保つ")
item("トラブル時の対応に必要なログの収集を常に行う")
item("可能な場面では有線LANの利用も検討する")

# =================================================================== 3 まとめ
heading("3", "まとめ")
body("今回はインターネットの概要、およびそれに関わる通信機器と、それらの間で情報を媒介するネット"
     "ワークケーブルについて学んだ。実習では実際にUTPストレートケーブルとUTPクロスケーブルを作製し、"
     "ケーブルテスターでその接続を確認した。また、製作したUTPケーブルとブロードバンドルータを用いて"
     "有線ネットワークと無線ネットワークを構築し、pingによる接続性の確認や、無線ネットワークの"
     "セキュリティ設定を行う方法を学んだ。")
body("実際に作業してみて、RJ-45コネクタに8本の導線を規格どおりの色順で差し込む作業は思いのほか"
     "細かく、被覆の剥き加減やより戻しの長さに気を遣う必要があった。今回はケーブルテスターでの診断に"
     "問題は出ず一度で接続できたが、配線を1本でも間違えると通信できなくなるため、ケーブル製作には"
     "丁寧さと正確さが重要だと実感した。各ネットワークやケーブルの性質を理解したうえで、正しく安全に"
     "インターネットを使えるよう、今後も注意していきたい。")

# =================================================================== 参考文献
heading("参考文献", "", level=1)
refs = [
    "[1] 前期電気基礎実験 実験テキスト「11章 ネットワーク実習」、東京理科大学工学部電気工学科",
    "[2] 前期電気基礎実験 講義資料、東京理科大学工学部電気工学科",
    "[3] NTT西日本「インターネットとは？意味と仕組みをわかりやすく解説」　https://flets-w.com/chienetta/lifestyle/cb_internet19.html",
    "[4] 総務省「国民のためのサイバーセキュリティサイト」　https://www.soumu.go.jp/",
    "[5] Rekisiru、京藤一葉「インターネットの歴史とは？年表から見るネット社会の普及」　https://rekisiru.com/7439",
    "[6] 1分で読めるIT用語辞典、佐々木真「ファイバー・オプティック・ケーブル」　https://wa3.i-3-i.info/word18452.html",
    "[7] Panduit Corp「イーサネットとは？LANケーブルとの違いや通信規格、種類を解説」　https://www.panduit.co.jp/column/naruhodo/9237/",
    "[8] ELECOM「有線LANの基礎知識」　https://www.elecom.co.jp/pickup/column/wired_column/00003/",
    "[9] HP Development Company「企業で安全に無線LANを使おう！セキュリティ強化に必要な対策を徹底解説」　https://jp.ext.hp.com/techdevice/cybersecuritysc/43/",
    "[10] WDIC Creators club「ネットマスク（通信用語の基礎知識）」　https://www.wdic.org/",
    "[11] 日本ネットワークインフォメーションセンター、宇夫陽次朗「RFCってなに？」　https://www.nic.ad.jp/ja/rfc-jp/WhatisRFC.html",
    "[12] NTT Communications Corporation「ゲートウェイとは？ルータとの違いや設定方法を初心者にもわかりやすく解説」　https://www.ntt.com/business/services/network/internet-connect/ocn-business/bocn/knowledge/archive_104",
]
for r in refs:
    p = para(r, size=9.5, space_after=3)
    p.paragraph_format.left_indent = Mm(8); p.paragraph_format.first_line_indent = Mm(-8)

page_number_footer()

out = os.path.join(os.path.dirname(__file__), "ネットワーク実習_実験レポート.docx")
doc.save(out)
print("saved:", out)
