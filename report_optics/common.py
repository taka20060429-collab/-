# -*- coding: utf-8 -*-
"""物理学実験レポート（Word）作成の共通部品。
熱電対レポート（report_thermo）と同じ書式：黒字・明朝/Times・A4・図は下/表は上にキャプション・式番号。
グラフは Excel のグラフオブジェクト（Excel 既定の書式）として Word に埋め込む。"""
import hashlib
import io
import math
import os
import re
import zipfile

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager as fm
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook
from openpyxl.chart import ScatterChart
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from PIL import Image

JP_FONT = "/usr/share/fonts/opentype/ipafont-gothic/ipagp.ttf"
if os.path.exists(JP_FONT):
    fm.fontManager.addfont(JP_FONT)
    plt.rcParams["font.family"] = fm.FontProperties(fname=JP_FONT).get_name()
plt.rcParams["axes.unicode_minus"] = False
plt.rcParams["mathtext.fontset"] = "dejavuserif"

BODY, HEAD, LATIN = "ＭＳ 明朝", "ＭＳ ゴシック", "Times New Roman"
TODO = "【要記入】"
CENTER, LEFT = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT
TEXT_W = 160  # 本文幅 [mm]

doc = None
EQDIR = None
CNT = {}
PART = [""]          # 部の記号（"A"，"B"）．番号は A-1，A-2 … の形になる
_CHART_N = [0]


def new_doc(eqdir):
    global doc, EQDIR
    EQDIR = eqdir
    os.makedirs(EQDIR, exist_ok=True)
    CNT.update({"fig": 0, "tab": 0, "eq": 0})
    _CHART_N[0] = 0
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = LATIN
    st.font.size = Pt(10.5)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), BODY)
    st.paragraph_format.line_spacing = 1.2
    st.paragraph_format.space_after = Pt(0)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin = sec.bottom_margin = Mm(25)
    sec.left_margin = sec.right_margin = Mm(25)
    return doc


# ---------------------------------------------------------------- 文字
def font(run, latin=LATIN, ea=BODY, size=None, bold=False):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:ascii"), latin)
    rf.set(qn("w:hAnsi"), latin)
    rf.set(qn("w:eastAsia"), ea)
    if size:
        run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


_MK = re.compile(r"(_\{[^}]*\}|\^\{[^}]*\})")


def add_text(p, text, latin=LATIN, ea=BODY, size=None, bold=False):
    """_{..} を下付き，^{..} を上付きの文字として書き込む。"""
    for tok in _MK.split(text):
        if not tok:
            continue
        if tok[:2] in ("_{", "^{"):
            r = p.add_run(tok[2:-1])
            font(r, latin, ea, size, bold)
            if tok[0] == "_":
                r.font.subscript = True
            else:
                r.font.superscript = True
        else:
            font(p.add_run(tok), latin, ea, size, bold)


def para(text="", size=10.5, bold=False, align=None, ea=BODY, latin=LATIN, sb=0, sa=4,
         indent=None, left=None, line=1.35, keep=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(sb), Pt(sa), line
    if indent is not None:
        pf.first_line_indent = Pt(indent)
    if left is not None:
        pf.left_indent = Mm(left)
    if keep:
        pf.keep_with_next = True
    if text:
        add_text(p, text, latin, ea, size, bold)
    return p


def body(text):
    return para(text, indent=10.5)


def item(text, left=4):
    p = para(text, line=1.35, sa=3)
    p.paragraph_format.first_line_indent = Mm(-4)
    p.paragraph_format.left_indent = Mm(left + 4)
    return p


def heading(text, level=1):
    size = {1: 13, 2: 11.5, 3: 10.5}[level]
    p = para(sb={1: 14, 2: 9, 3: 6}[level], sa=5, keep=True)
    font(p.add_run(text), HEAD, HEAD, size, True)
    return p


def todo(text):
    p = para(left=4, sa=4)
    font(p.add_run(text), LATIN, BODY, 10.5, True)
    return p


def nxt(kind):
    CNT[kind] += 1
    return f"{PART[0]}-{CNT[kind]}" if PART[0] else CNT[kind]


def set_part(prefix):
    """部を切り替え，図・表・式の番号を 1 から付け直す。"""
    PART[0] = prefix
    CNT.update({"fig": 0, "tab": 0, "eq": 0})


# ---------------------------------------------------------------- 図・表・式
def figure(path, caption, num, width=120):
    p = para(align=CENTER, sb=6, sa=2, keep=True)
    p.add_run().add_picture(path, width=Mm(width))
    c = para(align=CENTER, sa=10)
    font(c.add_run(f"図{num}　{caption}"), LATIN, BODY, 9.5)


def table(caption, num, headers, rows, widths, size=9.5, aligns=None):
    c = para(align=CENTER, sb=6, sa=3, keep=True)
    font(c.add_run(f"表{num}　{caption}"), LATIN, BODY, 9.5)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ""
        q = cell.paragraphs[0]
        q.alignment = CENTER
        add_text(q, h, size=size)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = t.rows[i + 1].cells[j]
            cell.text = ""
            q = cell.paragraphs[0]
            q.alignment = aligns[j] if aligns else CENTER
            add_text(q, str(v), size=size)
    for j, w in enumerate(widths):
        t.columns[j].width = Mm(w)
    for r in t.rows:
        for j, w in enumerate(widths):
            r.cells[j].width = Mm(w)
    para(sa=6)
    return t


def _render_tex(tex, size):
    path = os.path.join(EQDIR, "eq" + hashlib.md5(f"{tex}|{size}".encode()).hexdigest()[:12] + ".png")
    if not os.path.exists(path):
        fig = plt.figure(figsize=(0.01, 0.01))
        fig.text(0, 0, f"${tex}$", fontsize=size)
        fig.savefig(path, dpi=600, bbox_inches="tight", pad_inches=0.03, transparent=True)
        plt.close(fig)
    w, _ = Image.open(path).size
    return path, w / 600 * 25.4


def eq(tex, numbered=True, size=12.5):
    path, wmm = _render_tex(tex, size)
    wmm = min(wmm, TEXT_W - 22)
    p = para(sb=3, sa=5, line=1.0)
    ts = p.paragraph_format.tab_stops
    ts.add_tab_stop(Mm(TEXT_W / 2), WD_TAB_ALIGNMENT.CENTER)
    ts.add_tab_stop(Mm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
    font(p.add_run("\t"))
    p.add_run().add_picture(path, width=Mm(wmm))
    if numbered:
        n = nxt("eq")
        font(p.add_run(f"\t({n})"))
        return n
    return None


SUP = str.maketrans("0123456789-", "⁰¹²³⁴⁵⁶⁷⁸⁹⁻")


def sci(x, sig=3):
    """本文用：6.33×10⁻⁷ のような表記（有効数字 sig 桁）。"""
    if x == 0:
        return "0"
    e = int(math.floor(math.log10(abs(x))))
    m = x / 10 ** e
    if round(abs(m), sig - 1) >= 10:
        m /= 10
        e += 1
    return f"{m:.{sig - 1}f}×10{str(e).translate(SUP)}"


def tsci(x, sig=3):
    """式画像用（mathtext）：6.33\\times10^{-7}。"""
    s = sci(x, sig)
    m, e = s.split("×10")
    return f"{m}\\times 10^{{{e.translate(str.maketrans('⁰¹²³⁴⁵⁶⁷⁸⁹⁻', '0123456789-'))}}}"


# ---------------------------------------------------------------- 表紙・参考文献・保存
def cover(meta, title):
    for _ in range(4):
        para(sa=4)
    para(f"第 {meta.get('kai') or '　'} 回　物理学実験　報告書", size=16, bold=True, align=CENTER,
         ea=HEAD, latin=HEAD, sa=26)
    para(f"実験題目：{title}", size=15, bold=True, align=CENTER, ea=HEAD, latin=HEAD, sa=40)
    for s in [f"実験年月日：{meta.get('date') or TODO}", f"報告書提出日：{meta.get('submit') or TODO}",
              f"班名：{meta.get('han') or TODO}",
              f"実験者氏名：{meta['name']}　　学年：{meta['grade']}　　学籍番号：{meta['student_id']}"]:
        para(s, size=11.5, left=28, sa=10)
    para("共同実験者：", size=11.5, left=28, sa=6)
    for nm, sid in meta.get("coworkers") or [(TODO, TODO)]:
        para(f"氏名：{nm}　　学籍番号：{sid}", size=11, left=38, sa=4)
    doc.add_page_break()


def references(refs):
    heading("参考文献")
    for r in refs:
        p = para(r, size=10, sa=3, line=1.3)
        p.paragraph_format.left_indent = Mm(7)
        p.paragraph_format.first_line_indent = Mm(-7)


def save(path):
    sec = doc.sections[0]
    fp = sec.footer.paragraphs[0]
    fp.alignment = CENTER
    run = fp.add_run()
    for tag, attr, val in [("w:fldChar", "w:fldCharType", "begin"), ("w:instrText", None, "PAGE"),
                           ("w:fldChar", "w:fldCharType", "end")]:
        el = OxmlElement(tag)
        if attr:
            el.set(qn(attr), val)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = val
        run._r.append(el)
    font(run, size=10)
    sec.different_first_page_header_footer = True
    doc.save(path)


# ================================================================ Excel グラフ（Excel 既定の書式）
BLUE, ORANGE, GRAY, YELLOW = "4472C4", "ED7D31", "A5A5A5", "FFC000"
AXIS_GRAY, TEXT_GRAY = "D9D9D9", "595959"
C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def marker(sym, size, color):
    return Marker(symbol=sym, size=size,
                  spPr=GraphicalProperties(solidFill=color, ln=LineProperties(solidFill=color, w=9525)))


def line(series, color=None, width=28575, dash=None, none=False):
    gp = series.graphicalProperties
    if none:
        gp.line.noFill = True
        return
    gp.line.solidFill = color
    gp.line.width = width
    if dash:
        gp.line.dashStyle = dash


def base_chart(xtitle, ytitle, w=16, h=9.5):
    ch = ScatterChart()
    ch.style = 2
    ch.title = None
    ch.x_axis.title = xtitle
    ch.y_axis.title = ytitle
    for ax in (ch.x_axis, ch.y_axis):
        ax.delete = False
        ax.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill=AXIS_GRAY, w=9525)))
        ax.majorTickMark = "none"
        ax.minorTickMark = "none"
        ax.spPr = GraphicalProperties(ln=LineProperties(solidFill=AXIS_GRAY, w=9525))
        ax.number_format = "General"
    ch.legend.position = "b"
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.width, ch.height = w, h
    return ch


def _q(t):
    return f"{{{C_NS}}}{t}"


def _fix_layout(root):
    """Excel が保存するときと同じ指定：題名・凡例を重ねない，タイトルなし，游ゴシック 9 pt。"""
    if root.find(_q("roundedCorners")) is None:
        rc = etree.Element(_q("roundedCorners"))
        rc.set("val", "0")
        root.insert(0, rc)
    for title in root.iter(_q("title")):
        if title.find(_q("overlay")) is None:
            ov = etree.Element(_q("overlay"))
            ov.set("val", "0")
            title.find(_q("tx")).addnext(ov)
    for legend in root.iter(_q("legend")):
        if legend.find(_q("overlay")) is None:
            ov = etree.Element(_q("overlay"))
            ov.set("val", "0")
            legend.find(_q("legendPos")).addnext(ov)
    chart = root.find(_q("chart"))
    if chart.find(_q("autoTitleDeleted")) is None:
        atd = etree.Element(_q("autoTitleDeleted"))
        atd.set("val", "1")
        chart.insert(0, atd)
    if root.find(_q("txPr")) is None:
        tx = etree.fromstring(
            f'<c:txPr xmlns:c="{C_NS}" xmlns:a="{A_NS}"><a:bodyPr/><a:lstStyle/><a:p><a:pPr>'
            f'<a:defRPr sz="900"><a:solidFill><a:srgbClr val="{TEXT_GRAY}"/></a:solidFill>'
            f'<a:latin typeface="游ゴシック"/><a:ea typeface="游ゴシック"/></a:defRPr></a:pPr>'
            f'<a:endParaRPr lang="ja-JP"/></a:p></c:txPr>')
        sppr = root.find(_q("spPr"))
        (sppr if sppr is not None else chart).addnext(tx)
    return root


def _patch_xlsx(blob):
    zin = zipfile.ZipFile(io.BytesIO(blob))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if re.match(r"xl/charts/chart\d+\.xml$", it.filename):
                data = etree.tostring(_fix_layout(etree.fromstring(data)), xml_declaration=True,
                                      encoding="UTF-8", standalone=True)
            zout.writestr(it, data)
    return out.getvalue()


def _ref_values(ws, formula):
    m = re.match(r"^'?(.*?)'?!\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+)$", formula)
    col, r0, r1 = m.group(2), int(m.group(3)), int(m.group(5))
    return [ws[f"{col}{r}"].value for r in range(r0, r1 + 1)]


def native(ws, ch, anchor="H2"):
    """ワークシート ws とそのデータを参照するグラフ ch から，Word 埋め込み用の部品を作る。"""
    size = (ch.width, ch.height)
    ws.add_chart(ch, anchor)
    bio = io.BytesIO()
    ws.parent.save(bio)
    xlsx = _patch_xlsx(bio.getvalue())
    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(xlsx)).read("xl/charts/chart1.xml"))
    for ref in root.iter(_q("numRef")):
        vals = _ref_values(ws, ref.find(_q("f")).text)
        cache = etree.SubElement(ref, _q("numCache"))
        etree.SubElement(cache, _q("formatCode")).text = "General"
        etree.SubElement(cache, _q("ptCount")).set("val", str(len(vals)))
        for i, v in enumerate(vals):
            if isinstance(v, (int, float)):
                pt = etree.SubElement(cache, _q("pt"))
                pt.set("idx", str(i))
                etree.SubElement(pt, _q("v")).text = repr(float(v))
    ext = etree.SubElement(root, _q("externalData"))
    ext.set(f"{{{R_NS}}}id", "rId1")
    etree.SubElement(ext, _q("autoUpdate")).set("val", "0")
    return xlsx, etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True), size


def chart_figure(nat, caption, num, width=125):
    """Excel のグラフ（チャート＋埋め込みブック）を図として挿入する。"""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    xlsx, xml, (wcm, hcm) = nat
    _CHART_N[0] += 1
    n = _CHART_N[0]
    pkg = doc.part.package
    cp = Part(PackURI(f"/word/charts/chart{n}.xml"),
              "application/vnd.openxmlformats-officedocument.drawingml.chart+xml", b"", pkg)
    xp = Part(PackURI(f"/word/embeddings/Microsoft_Excel_Worksheet{n}.xlsx"),
              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", xlsx, pkg)
    rid_x = cp.relate_to(xp, RT.PACKAGE)
    cp._blob = xml.replace(b'r:id="rId1"', f'r:id="{rid_x}"'.encode())
    rid_c = doc.part.relate_to(cp, RT.CHART)
    cx, cy = int(width * 36000), int(width * hcm / wcm * 36000)
    p = para(align=CENTER, sb=6, sa=2, keep=True)
    p.add_run()._r.append(parse_xml(
        '<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        f'xmlns:a="{A_NS}" xmlns:c="{C_NS}" xmlns:r="{R_NS}">'
        f'<wp:inline distT="0" distB="0" distL="0" distR="0"><wp:extent cx="{cx}" cy="{cy}"/>'
        f'<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:docPr id="{5000 + n}" name="グラフ {n}"/>'
        '<wp:cNvGraphicFramePr/><a:graphic><a:graphicData '
        f'uri="{C_NS}"><c:chart r:id="{rid_c}"/></a:graphicData></a:graphic></wp:inline></w:drawing>'))
    c = para(align=CENTER, sa=10)
    font(c.add_run(f"図{num}　{caption}"), LATIN, BODY, 9.5)


def new_sheet(title):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    return ws
