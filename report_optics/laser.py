# -*- coding: utf-8 -*-
"""6. レーザ光の回折と干渉の実験 — 報告書（Word）生成。 python3 laser.py"""
import math
import os

import matplotlib.pyplot as plt
import numpy as np
from openpyxl.chart import Reference, Series
from openpyxl.chart.marker import Marker

import common as C
from common import body, eq, heading, item, nxt, para, sci, table, tsci

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "img_laser")
os.makedirs(IMG, exist_ok=True)

META = {"kai": None, "date": None, "submit": None, "han": "D班", "name": "森木崇允", "grade": "1",
        "student_id": "4326091",
        "coworkers": [("石井真人", "4326007"), ("木元喜紹", "4326029"), ("陶山稜平", "4326051"),
                      ("林田知子", "4326071")]}

# ---------------------------------------------------------------- 測定値（実験ノートの Excel より）
LAM = 632.8e-9     # He-Ne レーザの波長 [m]
L_SLIT = 1.40      # 単スリットとスクリーンの距離 l [m]
L_PIN = 1.59       # ピンホールとスクリーンの距離 [m]
L_DBL = 1.59       # 複スリットとスクリーンの距離 [m]
L_GR = 1.59        # 回折格子とスクリーンの距離 [m]
SLIT = [("狭いスリット", 29.6, 0.06), ("広いスリット", 16.9, 0.12)]        # (名前, PQ=x1 [mm], 公称の幅 [mm])
PIN = [(0.1, 14.70), (0.2, 12.15), (0.3, 10.00)]                          # (公称の直径 r [mm], R1 [mm])
DOUBLE = [(0.4, 5, 3.15), (0.2, 2, 3.95)]    # (公称の間隔 d [mm], 測定した干渉じまの数, 暗線の間隔Δx [mm])
GRATING_X = 202.0                            # 1次の光点の間隔 PP′ [mm]
GRATING_NOM = 1000                           # 回折格子の公称の本数 [本/cm]
READ = 0.5                                   # スクリーン上の長さの読み取り誤差 [mm]（目安）

# ---------------------------------------------------------------- 計算
slit = []
for name, x, nom in SLIT:
    th = x * 1e-3 / (2 * L_SLIT)
    d = LAM / th
    slit.append(dict(name=name, x=x, nom=nom, th=th, d=d * 1e3, err=(d * 1e3 - nom) / nom * 100,
                     rel_read=READ / x * 100))
pin = []
for nom, R in PIN:
    th = R * 1e-3 / (2 * L_PIN)
    r = 1.22 * LAM / th
    pin.append(dict(nom=nom, R=R, th=th, r=r * 1e3, err=(r * 1e3 - nom) / nom * 100, rel_read=READ / R * 100,
                    R_th=2.44 * L_PIN * LAM / (nom * 1e-3) * 1e3))
dbl = []
for nom, k, dx in DOUBLE:
    d = L_DBL * LAM / (dx * 1e-3)
    dbl.append(dict(nom=nom, k=k, dx=dx, d=d * 1e3, err=(d * 1e3 - nom) / nom * 100,
                    rel_read=READ / dx * 100, rel_read_k=READ / (dx * k) * 100,
                    dx_th=L_DBL * LAM / (nom * 1e-3) * 1e3))
gth = GRATING_X * 1e-3 / (2 * L_GR)
g_d = LAM / gth                         # [m]
g_n = 1 / (g_d * 100)                   # 1 cm あたりの本数
g_err = (g_n - GRATING_NOM) / GRATING_NOM * 100
g_d_exact = LAM / math.sin(math.atan(gth))
th_max = gth
beam10 = 0.7 + 10 * 1.2e-3 * 1e3      # [mm]
ALL = ([("単スリット（" + s_["name"] + "）", s_["err"], s_["rel_read"]) for s_ in slit] +
       [(f"ピンホール（{p_['nom']:.1f} mm）", p_["err"], p_["rel_read"]) for p_ in pin] +
       [(f"複スリット（{d_['nom']:.1f} mm）", d_["err"], d_["rel_read"]) for d_ in dbl])


# ---------------------------------------------------------------- 模式図（自作）
def fig_single(path):
    fig, ax = plt.subplots(figsize=(6.2, 2.7))
    ax.plot([0, 0], [-1.6, -0.18], color="black", lw=2.2)
    ax.plot([0, 0], [0.18, 1.6], color="black", lw=2.2)
    ax.plot([7, 7], [-1.7, 1.7], color="black", lw=1.4)
    for y in (-0.6, -0.3, 0.0, 0.3, 0.6):
        ax.annotate("", xy=(-0.05, y * 0.25), xytext=(-1.6, y * 0.25), arrowprops=dict(arrowstyle="-|>", lw=0.9))
    ax.text(-1.6, 0.35, "レーザ光", fontsize=8.5)
    ax.axhline(0, color="0.6", lw=0.6, ls=":")
    ax.plot([0, 7], [0, 1.0], color="black", lw=0.9)
    ax.plot([0, 7], [0, -1.0], color="black", lw=0.9)
    ax.plot([7], [1.0], "o", color="black", ms=3); ax.plot([7], [-1.0], "o", color="black", ms=3)
    ax.text(7.15, 1.0, "P（第1暗点）", va="center", fontsize=8.5)
    ax.text(7.15, -1.0, "Q（第1暗点）", va="center", fontsize=8.5)
    ax.annotate("", xy=(7.9, 1.0), xytext=(7.9, -1.0), arrowprops=dict(arrowstyle="<->", lw=0.8))
    ax.text(8.0, 0, "x＝PQ", va="center", fontsize=9)
    from matplotlib.patches import Arc
    ax.add_patch(Arc((0, 0), 3.0, 3.0, theta1=0, theta2=math.degrees(math.atan(1 / 7)), lw=0.8))
    ax.text(1.65, 0.08, r"$\theta_1$", fontsize=10)
    ax.annotate("", xy=(7, -1.45), xytext=(0, -1.45), arrowprops=dict(arrowstyle="<->", lw=0.8))
    ax.text(3.5, -1.4, "l", ha="center", va="bottom", fontsize=10, style="italic")
    ax.text(0, 1.75, "単スリット（幅 d）／ピンホール（直径 r）", ha="center", fontsize=8.5)
    ax.text(7, 1.85, "スクリーン", ha="center", fontsize=8.5)
    ax.set_xlim(-1.8, 9.6); ax.set_ylim(-1.8, 2.1); ax.axis("off")
    fig.tight_layout(); fig.savefig(path, dpi=220); plt.close(fig)


def fig_double(path):
    fig, ax = plt.subplots(figsize=(6.2, 2.7))
    ax.plot([0, 0], [-1.6, -0.45], color="black", lw=2.2)
    ax.plot([0, 0], [-0.25, 0.25], color="black", lw=2.2)
    ax.plot([0, 0], [0.45, 1.6], color="black", lw=2.2)
    ax.plot([7, 7], [-1.7, 1.7], color="black", lw=1.4)
    ax.text(-0.25, 0.35, "A", ha="right", va="center", fontsize=9)
    ax.text(-0.25, -0.35, "B", ha="right", va="center", fontsize=9)
    ax.axhline(0, color="0.6", lw=0.6, ls=":")
    ax.plot([0, 7], [0.35, 1.1], color="black", lw=0.9)
    ax.plot([0, 7], [-0.35, 1.1], color="black", lw=0.9)
    ax.plot([7], [1.1], "o", color="black", ms=3)
    ax.text(7.15, 1.1, "P", va="center", fontsize=9)
    ax.text(7.15, 0, "O", va="center", fontsize=9)
    ax.annotate("", xy=(7.6, 1.1), xytext=(7.6, 0), arrowprops=dict(arrowstyle="<->", lw=0.8))
    ax.text(7.7, 0.55, "x", va="center", fontsize=10, style="italic")
    ax.annotate("", xy=(-0.6, 0.35), xytext=(-0.6, -0.35), arrowprops=dict(arrowstyle="<->", lw=0.8))
    ax.text(-0.75, 0, "d", ha="right", va="center", fontsize=10, style="italic")
    ax.annotate("", xy=(7, -1.45), xytext=(0, -1.45), arrowprops=dict(arrowstyle="<->", lw=0.8))
    ax.text(3.5, -1.4, "l", ha="center", va="bottom", fontsize=10, style="italic")
    ax.text(0, 1.75, "複スリット", ha="center", fontsize=8.5)
    ax.text(7, 1.85, "スクリーン", ha="center", fontsize=8.5)
    ax.text(3.8, 0.25, "光路差 BP－AP ≈ xd / l", fontsize=8.5)
    ax.set_xlim(-1.8, 9.0); ax.set_ylim(-1.8, 2.1); ax.axis("off")
    fig.tight_layout(); fig.savefig(path, dpi=220); plt.close(fig)


# ---------------------------------------------------------------- Excel グラフ
def chart_nominal():
    ws = C.new_sheet("比較")
    ws.append(["公称値 [mm]", "単スリットの幅 [mm]", "ピンホールの直径 [mm]", "複スリットの間隔 [mm]", "", "x", "y = x"])
    rows = [(s["nom"], s["d"], None, None) for s in slit] + [(p["nom"], None, p["r"], None) for p in pin] + \
           [(d["nom"], None, None, d["d"]) for d in dbl]
    for r in rows:
        ws.append(list(r))
    ws["F2"], ws["G2"], ws["F3"], ws["G3"] = 0, 0, 0.45, 0.45
    n = len(rows) + 1
    ch = C.base_chart("公称値 [mm]", "測定値 [mm]", h=10)
    s0 = Series(Reference(ws, min_col=7, min_row=2, max_row=3), Reference(ws, min_col=6, min_row=2, max_row=3),
                title="測定値＝公称値")
    s0.marker = Marker(symbol="none"); C.line(s0, C.GRAY, width=19050, dash="dash")
    ch.series.append(s0)
    for col, color, sym, name in [(2, C.BLUE, "circle", "単スリットの幅"), (3, C.ORANGE, "square", "ピンホールの直径"),
                                  (4, C.YELLOW, "triangle", "複スリットの間隔")]:
        s = Series(Reference(ws, min_col=col, min_row=2, max_row=n), Reference(ws, min_col=1, min_row=2, max_row=n),
                   title=name)
        s.marker = C.marker(sym, 8, color); C.line(s, none=True)
        ch.series.append(s)
    for ax in (ch.x_axis, ch.y_axis):
        ax.scaling.min, ax.scaling.max, ax.majorUnit = 0, 0.45, 0.05
    return C.native(ws, ch, "I2")


def chart_pinhole():
    ws = C.new_sheet("ピンホール")
    ws.append(["1/r [mm⁻¹]", "R1 測定 [mm]", "", "1/r [mm⁻¹]", "R1 理論 [mm]"])
    for p in pin:
        ws.append([1 / p["nom"], p["R"]])
    k = 2.44 * L_PIN * LAM * 1e6      # R1 [mm] = k / r [mm]
    for i, x in enumerate(np.linspace(0, 11, 12)):
        ws.cell(row=2 + i, column=4, value=float(x))
        ws.cell(row=2 + i, column=5, value=float(k * x))
    ch = C.base_chart("1/r [mm⁻¹]（r：ピンホールの公称の直径）", "第一暗環の直径 R₁ [mm]", h=9.5)
    s_t = Series(Reference(ws, min_col=5, min_row=2, max_row=13), Reference(ws, min_col=4, min_row=2, max_row=13),
                 title="理論 R₁＝2.44 lλ / r")
    s_t.marker = Marker(symbol="none"); C.line(s_t, C.ORANGE)
    s_m = Series(Reference(ws, min_col=2, min_row=2, max_row=4), Reference(ws, min_col=1, min_row=2, max_row=4),
                 title="測定値")
    s_m.marker = C.marker("circle", 8, C.BLUE); C.line(s_m, none=True)
    ch.series.append(s_t); ch.series.append(s_m)
    ch.x_axis.scaling.min, ch.x_axis.scaling.max, ch.x_axis.majorUnit = 0, 11, 1
    ch.y_axis.scaling.min, ch.y_axis.scaling.max, ch.y_axis.majorUnit = 0, 25, 5
    return C.native(ws, ch, "G2"), k


# ================================================================ 文書
STANDALONE = __name__ == "__main__"
if STANDALONE:
    C.new_doc(os.path.join(IMG, "eq"))
fig_single(os.path.join(IMG, "fig_single.png"))
fig_double(os.path.join(IMG, "fig_double.png"))
if STANDALONE:
    C.cover(META, "レーザ光の回折と干渉の実験")
heading("B．レーザー光の回折と干渉の実験")

heading("B.1　目的")
body("He-Neレーザの光を用いて，単スリットとピンホール（小円孔）による回折，複スリットと回折格子による干渉を"
     "観測する．スクリーン上に現れる回折じま・干渉じまの間隔から，スリットの幅，ピンホールの直径，スリットの"
     "間隔および回折格子の格子定数を求め，公称値と比較することで，光の波としての性質を確かめる．")

heading("B.2　原理")
heading("(1)　単スリットによる回折", 2)
f_s = nxt("fig")
body("幅dの細いスリットに波長λの光を当てると，光はスリットの縁で回り込み（回折），スクリーン上に明暗のしまが"
     "現れる．スリット内の各点から出た光の位相のずれを考えると，入射方向から角θの向きでは，スリットの両端を"
     "通る光の光路差が d sinθ になる．この光路差がλの整数倍のとき，スリットをいくつかに等分した部分どうしの"
     "光が打ち消し合って暗くなる[1][2]．すなわち暗線の条件は")
n_dark = eq(r"d\sin\theta=n\lambda\qquad(n=1,2,3,\ldots)")
body(f"であり，中央の明るい帯の両側にある最初の暗点（n＝1）を P，Q とする（図{f_s}）．スクリーンまでの距離 l "
     "が十分大きいとき，PQ＝x として θ_{1}≈sinθ_{1}≈tanθ_{1}＝x/(2l) と近似できるので，")
n_slit = eq(r"d=\dfrac{\lambda}{\sin\theta_1}\approx\dfrac{2l\lambda}{x}")
body("からスリットの幅dが求まる．")
C.figure(os.path.join(IMG, "fig_single.png"), "単スリット（ピンホール）による回折の配置（模式図）", f_s, 120)

heading("(2)　ピンホールによる回折", 2)
body("直径rの円形の穴（ピンホール）による回折では，同心円状のしまが現れる．いちばん内側の暗い円（第一暗環）の"
     "直径をR_{1}，対応する回折角をθ_{1}とすると，円形の開口に対する計算から sinθ_{1}＝1.220λ/r であることが"
     "知られている（第二，第三暗環ではそれぞれ 2.232λ/r，3.238λ/r）[1][2]．単スリットと同様に "
     "θ_{1}≈R_{1}/(2l) と近似すると，")
n_pin = eq(r"r=\dfrac{1.22\lambda}{\sin\theta_1}\approx\dfrac{2.44\,l\lambda}{R_1}")
body("となる．係数が1でなく1.22になるのは，円形の開口では中心付近ほど開口の幅が広く，回折への寄与の仕方が"
     "スリットと異なるためである．")

heading("(3)　複スリットによる干渉（ヤングの実験）", 2)
f_d = nxt("fig")
body(f"間隔dの2本のスリットA，Bを通った光は，スクリーン上で重なり合って干渉する（図{f_d}）．スクリーン中央Oから"
     "距離xの点Pでの光路差は，x≪l，d≪l のとき BP－AP≈xd/l となる．光路差が波長の整数倍なら強め合って明線，"
     "半波長の奇数倍なら弱め合って暗線になるので，明線の位置は")
n_br = eq(r"x=\pm\dfrac{nl\lambda}{d}")
body("，暗線の位置は")
n_bk = eq(r"x=\pm\dfrac{(n+\frac{1}{2})\,l\lambda}{d}")
body("（n＝0, 1, 2, …）である[1]．隣り合う明線（または暗線）の間隔Δxは一定で，")
n_dx = eq(r"\Delta x=\dfrac{l\lambda}{d}\qquad\therefore\quad d=\dfrac{l\lambda}{\Delta x}")
body("からスリットの間隔dが求まる．")
C.figure(os.path.join(IMG, "fig_double.png"), "複スリットによる干渉の配置（模式図）", f_d, 120)

heading("(4)　回折格子による干渉", 2)
body("回折格子は，ガラス板に等間隔で多数の細い溝を刻んだもので，溝の間がスリットの役割をする．格子定数（溝の"
     "間隔）をdとすると，隣り合うスリットからの光の光路差 d sinθ が波長の整数倍になる方向で光が強め合う．")
n_gr = eq(r"d\sin\theta=n\lambda\qquad(n=0,1,2,\ldots)")
body(f"スリットの数が非常に多いため，この条件からわずかにずれた方向では多数の光が打ち消し合い，複スリットよりも"
     f"鋭い光点になる[1]．スクリーン上の1次の光点P，P′の間隔をxとすると，図{f_s}と同様に sinθ_{{1}}≈x/(2l) なので，")
n_grd = eq(r"d=\dfrac{\lambda}{\sin\theta_1}\approx\dfrac{2l\lambda}{x}")
body("から格子定数dと，1 cmあたりの溝の本数1/dが求まる．")

heading("B.3　装置")
body("本実験で用いたHe-Neレーザは，ヘリウムとネオンの混合気体を放電で励起し，2枚の反射鏡の間を往復する光を"
     "誘導放出によって増幅して取り出す装置である[1]．波長は632.8 nm（赤色），ビームの直径は約0.7 mm，ビームの"
     "広がり角は約1.2 mradであり，単色で位相のそろった（コヒーレントな）光が得られるため，回折や干渉のしまを"
     "鮮明に観測できる．")

# ---------------------------------------------------------------- 3 実験方法
heading("B.4　実験方法")
body("レーザ装置のアースが接続されていることを確認して電源を入れた．レーザ光を直接目に入れないこと，反射光を"
     "見つめないこと，スリットや回折格子に指で触れないことに注意した．光学台上で，スリット（またはピンホール，"
     f"回折格子）とスクリーンの距離 l は，単スリットでは {L_SLIT:.2f} m，ピンホール・複スリット・回折格子では"
     f" {L_PIN:.2f} m とした．")
heading("(1)　単スリットによる回折", 2)
body("フィルムに幅の異なる2つの窓をあけた単スリットを光学台のマウントに取り付け，はじめに狭い方，次に広い方に"
     "レーザ光を当てた．スクリーン上の回折じまを見ながらスリットとレーザの位置を調整し，中央の明るい帯をはさむ"
     "第1暗点P，Qの間隔 x＝PQ を測定した．式(" + str(n_slit) + ")からスリットの幅dを求めた．")
heading("(2)　ピンホールによる回折", 2)
body("直径0.1 mm，0.2 mm，0.3 mmの3種類のピンホールに交換し，スクリーン上の同心円状のしまのうち第一暗環の直径"
     "R_{1}を測定した．式(" + str(n_pin) + ")からピンホールの直径rを求めた．")
heading("(3)　複スリットによる干渉", 2)
body("スリットの間隔が d＝0.4 mm と 0.2 mm の2種類の複スリット（スリットの幅はいずれも0.1 mm）に交換し，"
     "スクリーン上の干渉じまについて，隣り合う暗線の間隔Δxを測定した．式(" + str(n_dx) + ")からdを求めた．")
heading("(4)　回折格子による干渉", 2)
body("回折格子に交換し，スクリーン上の0次の光点の左右に現れる1次の光点P，P′の間隔xを測定した．式("
     + str(n_grd) + ")から格子定数dと1 cmあたりの本数1/dを求めた．")

heading("(5)　使用器具", 2)
t_eq = nxt("tab")
body(f"使用した器具を表{t_eq}に示す．")
table("使用器具", t_eq, ["器具", "数量", "備考"],
      [["He-Neレーザ装置（島津理化 GLG-5003形）", "1台", "波長632.8 nm，出力0.5 mW以上"],
       ["レーザ用光学台（LD-43形）", "1台", "マウント用光学板，キャリアを含む"],
       ["単スリット（フィルム）", "1枚", "幅の異なる2つの窓"],
       ["ピンホール（フィルム）", "3種", "直径0.1，0.2，0.3 mm"],
       ["複スリット（フィルム）", "2種", "間隔0.2，0.4 mm，幅0.1 mm"],
       ["回折格子", "1枚", ""], ["スクリーン", "1枚", ""], ["巻尺，定規", "一式", "l とスクリーン上の長さの測定"]],
      [62, 18, 60], aligns=[C.LEFT, C.CENTER, C.LEFT])

# ---------------------------------------------------------------- 5 実験結果
heading("B.5　実験結果と考察")
body(f"レーザの波長は λ＝632.8 nm＝6.328×10⁻⁷ m である．スクリーンまでの距離は，単スリットで l＝{L_SLIT:.2f} m，"
     f"ピンホール・複スリット・回折格子で l＝{L_PIN:.2f} m である．")
heading("(1)　単スリットによる回折", 2)
t1 = nxt("tab")
body(f"測定した第1暗点の間隔xと，式(" + str(n_slit) + f")から求めたスリットの幅dを表{t1}に示す．")
table("単スリットの幅", t1, ["スリット", "x＝PQ [mm]", "θ_{1} [rad]", "d [mm]", "公称値 [mm]", "公称値との差 [%]"],
      [[s["name"], f"{s['x']:.1f}", sci(s["th"], 3), f"{s['d']:.3f}", f"{s['nom']:.2f}", f"{s['err']:+.1f}"]
       for s in slit], [30, 22, 26, 22, 24, 30])
s = slit[0]
body("狭いスリットを例に計算過程を示す．")
eq(rf"d\approx\dfrac{{2l\lambda}}{{x}}=\dfrac{{2\times{L_SLIT:.2f}\times6.328\times10^{{-7}}}}{{{s['x'] * 1e-3:.4f}}}"
   rf"={tsci(s['d'] * 1e-3, 3)}\ \mathrm{{m}}={s['d']:.3f}\ \mathrm{{mm}}", numbered=False, size=11)

heading("(2)　ピンホールによる回折", 2)
t2 = nxt("tab")
body(f"第一暗環の直径R_{{1}}と，式(" + str(n_pin) + f")から求めたピンホールの直径rを表{t2}に示す．")
table("ピンホールの直径", t2, ["公称の直径 [mm]", "R_{1} [mm]", "θ_{1} [rad]", "r [mm]", "公称値との差 [%]"],
      [[f"{p['nom']:.1f}", f"{p['R']:.2f}", sci(p["th"], 3), f"{p['r']:.3f}", f"{p['err']:+.1f}"] for p in pin],
      [30, 24, 26, 24, 32])
p = pin[0]
body("直径0.1 mmのピンホールを例に計算過程を示す．")
eq(rf"r\approx\dfrac{{2.44\,l\lambda}}{{R_1}}=\dfrac{{2.44\times{L_PIN:.2f}\times6.328\times10^{{-7}}}}"
   rf"{{{p['R'] * 1e-3:.5f}}}={tsci(p['r'] * 1e-3, 3)}\ \mathrm{{m}}={p['r']:.3f}\ \mathrm{{mm}}",
   numbered=False, size=11)

heading("(3)　複スリットによる干渉", 2)
t3 = nxt("tab")
body(f"測定した暗線の間隔Δxと，式(" + str(n_dx) + f")から求めたdを表{t3}に示す．参考として，公称の間隔から"
     f"計算されるΔxの値も併記した．")
table("複スリットの間隔", t3, ["公称の間隔 [mm]", "Δx 測定 [mm]", "Δx 公称値から [mm]", "d [mm]", "公称値との差 [%]"],
      [[f"{d['nom']:.1f}", f"{d['dx']:.2f}", f"{d['dx_th']:.2f}", f"{d['d']:.3f}", f"{d['err']:+.0f}"]
       for d in dbl], [28, 28, 34, 24, 32])
d = dbl[0]
body("間隔0.4 mmの複スリットを例に計算過程を示す．")
eq(rf"d=\dfrac{{l\lambda}}{{\Delta x}}=\dfrac{{{L_DBL:.2f}\times6.328\times10^{{-7}}}}{{{d['dx'] * 1e-3:.5f}}}"
   rf"={tsci(d['d'] * 1e-3, 3)}\ \mathrm{{m}}={d['d']:.3f}\ \mathrm{{mm}}", numbered=False, size=11)

heading("(4)　回折格子による干渉", 2)
body(f"1次の光点P，P′の間隔は x＝{GRATING_X:.1f} mm であった．式(" + str(n_grd) + ")より，")
eq(rf"d\approx\dfrac{{2l\lambda}}{{x}}=\dfrac{{2\times{L_GR:.2f}\times6.328\times10^{{-7}}}}{{{GRATING_X * 1e-3:.4f}}}"
   rf"={tsci(g_d, 3)}\ \mathrm{{m}}={tsci(g_d * 100, 3)}\ \mathrm{{cm}}", numbered=False, size=11)
body(f"となり，1 cmあたりの溝の本数は 1/d＝{g_n:.0f} 本/cm であった．公称値 {GRATING_NOM} 本/cm に対する相対誤差は"
     f"{abs(g_n - GRATING_NOM) / GRATING_NOM * 100:.1f} %である．")

f_c = nxt("fig")
body(f"(1)〜(3)で求めた値を公称値と比べて図{f_c}に示す．破線は測定値と公称値が等しい場合を表す．")
C.chart_figure(chart_nominal(), "測定値と公称値の比較", f_c, 115)
(nat_pin, k_pin) = chart_pinhole()
f_p = nxt("fig")
body(f"また，式(" + str(n_pin) + f")によれば第一暗環の直径R_{{1}}はピンホールの直径rに反比例する．測定したR_{{1}}を"
     f"公称の直径の逆数1/rに対して図{f_p}に示す．直線は R_{{1}}＝2.44lλ/r の理論値である．")
C.chart_figure(nat_pin, "ピンホールの直径の逆数と第一暗環の直径の関係", f_p, 115)

t5 = nxt("tab")
body(f"以上の結果を表{t5}にまとめる．")
table("実験結果のまとめ", t5, ["測定対象", "測定値", "公称値"],
      [[f"単スリット（{s['name']}）の幅", f"{s['d']:.3f} mm", f"{s['nom']:.2f} mm"] for s in slit] +
      [[f"ピンホールの直径（{p['nom']:.1f} mm）", f"{p['r']:.3f} mm", f"{p['nom']:.1f} mm"] for p in pin] +
      [[f"複スリットの間隔（{d['nom']:.1f} mm）", f"{d['d']:.3f} mm", f"{d['nom']:.1f} mm"] for d in dbl] +
      [["回折格子の格子定数", f"{sci(g_d * 100, 3)} cm（{g_n:.0f} 本/cm）", f"{GRATING_NOM} 本/cm"]],
      [62, 50, 30], aligns=[C.LEFT, C.CENTER, C.CENTER])

# ---------------------------------------------------------------- 6 考察
heading("(5)　考察", 2)
heading("(a)　測定値と公称値の比較", 3)
good = [n for n, e, _ in ALL if abs(e) < 5]
bad = [(n, e, rr) for n, e, rr in ALL if abs(e) >= 5]
n_pos = sum(1 for _, e, _ in ALL if e > 0)
n_neg = sum(1 for _, e, _ in ALL if e < 0)
body(f"図{f_c}より，測定値はいずれも公称値と同じ程度の大きさになり，光の波長とスクリーン上の長さだけから"
     f"0.1 mm程度の微小な寸法を求められることが確かめられた．公称値との差が5 %未満だったのは"
     + ("，".join(good) if good else "なし") + f"であり，回折格子も{g_err:+.1f} %と非常によく一致した．"
     f"一方で，" + "，".join(f"{n}（{e:+.0f} %）" for n, e, _ in bad) + "は差が大きかった．")
body(f"スクリーン上の長さの読み取り誤差を±{READ} mm とすると，それによる相対誤差は"
     + "，".join(f"{n}で{rr:.0f} %" for n, _, rr in bad) +
     "である．式(" + str(n_slit) + ")，(" + str(n_pin) + ")，(" + str(n_dx) + ")では求める寸法がスクリーン上の"
     "長さに反比例するので，長さの相対誤差がそのまま寸法の相対誤差になる．"
     + ("差の大きかった測定のうち，複スリットは読み取り誤差と同程度の差であり，間隔の小さいしまを1つ分だけ読んだ"
        "ことが主な原因と考えられる．" if all(abs(d_['err']) < 2.5 * d_['rel_read'] for d_ in dbl) and
        any(abs(d_['err']) >= 5 for d_ in dbl) else "")
     + f"また，公称値より大きい測定値が{n_pos}個，小さい測定値が{n_neg}個あり，"
     + ("ずれの向きがそろっていないことから，一方向の系統誤差よりも，ぼやけたしまの位置をどこと判断したかという"
        "読み取りのばらつきが主な原因と考えられる．" if n_pos and n_neg else
        "ずれの向きがそろっていることから，しまの位置を一方向にずらして読む系統的な癖があったと考えられる．"))
pin_sorted = sorted(pin, key=lambda p_: p_["nom"])
body(f"ピンホールでは，式(" + str(n_pin) + f")より第一暗環の直径R_{{1}}は直径rに反比例するはずである．図{f_p}を見ると，"
     f"rが大きいほどR_{{1}}は小さくなっており，反比例の傾向は確認できた．しかし理論値（{k_pin:.1f}/r mm）と比べると，"
     + "，".join(f"r＝{p_['nom']:.1f} mm では測定値 {p_['R']:.2f} mm に対し理論値 {p_['R_th']:.1f} mm"
                for p_ in pin_sorted)
     + f"であり，とくに直径0.1 mmのピンホールで差が大きい．小さいピンホールほど光が大きく広がって暗くなるため，"
     f"外側の暗環の位置を判断しにくく，内側の明るい部分の縁を暗環と見なしてR_{{1}}を小さく読んだ可能性がある．")
body(f"複スリットについて，公称の間隔から計算されるΔxは0.4 mmで{dbl[0]['dx_th']:.2f} mm，0.2 mmで"
     f"{dbl[1]['dx_th']:.2f} mm であり，測定値（{dbl[0]['dx']:.2f} mm，{dbl[1]['dx']:.2f} mm）との差は0.5 mm を"
     f"超えるものもある．暗線の間隔が数 mm と小さいため，1つ分の間隔を読むと読み取り誤差の影響が大きい．"
     f"5つ分の間隔をまとめて測れば，同じ読み取り誤差でも相対誤差は1/5になる"
     f"（0.4 mmのスリットで{dbl[0]['rel_read']:.0f} %→{dbl[0]['rel_read'] / 5:.0f} %）．")
body(f"回折格子では 1/d＝{g_n:.0f} 本/cm となり，公称値 {GRATING_NOM} 本/cm との差は{g_err:+.1f} %であった．"
     f"回折格子では光点が鋭く，しかも1次の光点の間隔が {GRATING_X:.0f} mm と長いため，±{READ} mm の読み取り誤差に"
     f"よる相対誤差は{READ / GRATING_X * 100:.1f} %にすぎない．他の測定に比べて格段に精度よく測定できたのはこのため"
     f"である．")

heading("(b)　小角近似の妥当性", 3)
body(f"本実験では sinθ≈tanθ≈θ と近似した．回折角が最も大きいのは回折格子の1次の光で，θ_{{1}}≈x/(2l)＝"
     f"{th_max:.4f} rad（約{math.degrees(th_max):.1f}°）である．近似を用いずに θ_{{1}}＝tan^{{-1}}(x/2l) から"
     f" d＝λ/sinθ_{{1}} を計算すると {sci(g_d_exact * 100, 4)} cm となり，近似した値との差は"
     f"{(g_d_exact - g_d) / g_d * 100:+.2f} %にすぎない．スリットやピンホールでは回折角がさらに小さい"
     f"（θ_{{1}}≈{min(x['th'] for x in slit + pin):.4f}〜{max(x['th'] for x in slit + pin):.4f} rad）ので，小角近似による誤差は"
     f"(a)で述べた読み取りの誤差に比べて無視できる．")

heading("B.6　課題")
para("(1)　コヒーレントな波とレーザ光の特徴", sa=2, keep=True)
body("コヒーレント（可干渉）な波とは，波の位相の関係が時間的・空間的に一定に保たれ，重ね合わせたときに安定した"
     "干渉じまをつくることができる波である．レーザ光は誘導放出によって位相のそろった光子が増えていくため，"
     "コヒーレンスが高い．その他の特徴として，(i) 単色性：波長の幅が非常に狭い，(ii) 指向性：ビームがほとんど広がら"
     "ずに直進する（本装置では広がり角約1.2 mrad），(iii) 高輝度：エネルギーを狭い範囲に集中できる，といった性質"
     "がある[1][2]．本実験で回折じま・干渉じまが鮮明に見えたのは，これらの性質による．")
para("(2)　偏光とその確かめ方", sa=2, keep=True)
body("光は進行方向に垂直に振動する横波であり，その電場の振動方向がそろっている光を偏光（直線偏光）という．"
     "振動方向がランダムな光は自然光（非偏光）とよばれる．偏光しているかどうかは，偏光板を光の通り道に置いて"
     "回転させ，透過光の明るさが変化するかで確かめられる．直線偏光なら，偏光板を回すと明るさが周期的に変わり，"
     "90°ごとに最大と最小（ほぼ0）になる（マリュスの法則）．本装置のように偏光がランダムな光では，偏光板を回しても"
     "明るさはほとんど変わらない[2]．")
para("(3)　回折の条件 d sinθ＝nλ の導出", sa=2, keep=True)
body("間隔dで並んだ2点（格子の隣り合うスリット，あるいは単スリットの両端）から，入射方向と角θをなす向きに"
     "進む光を考える．一方の点から他方の光線に垂線を下ろすと，2つの光線の道のりの差（光路差）は d sinθ となる．"
     "2つの光が強め合うのは光路差が波長の整数倍のときなので d sinθ＝nλ が得られる．単スリットの場合には，"
     "スリットを幅d/2ずつの2つの部分に分けると，対応する点どうしの光路差がλ/2になる方向（d sinθ＝λ）で全体が"
     "打ち消し合うので，同じ形の式が暗線の条件（式(" + str(n_dark) + ")）になる．")
para("(4)　ガスレーザと固体レーザの発振の仕組み", sa=2, keep=True)
body("レーザは，励起によって上の準位にある原子の数が下の準位より多い状態（反転分布）をつくり，誘導放出で光を"
     "増幅し，2枚の鏡からなる共振器で特定の波長の光を往復させて発振させる．He-Neレーザなどのガスレーザでは，"
     "放電で加速された電子がHe原子を励起し，そのエネルギーが衝突によってNe原子に移されて反転分布ができる．"
     "固体レーザ（ルビーレーザやNd:YAGレーザなど）では，結晶中に添加したイオンをフラッシュランプや半導体レーザの"
     "光で励起（光ポンピング）して反転分布をつくる[2]．")
para("(5)　液晶ディスプレイの原理", sa=2, keep=True)
body("液晶ディスプレイは，偏光の向きが直交した2枚の偏光板の間に液晶をはさんだ構造をもつ．電圧をかけないときは，"
     "ねじれて並んだ液晶分子に沿って光の偏光方向が90°回転するので，光は2枚目の偏光板を通過して明るくなる．"
     "電圧をかけると液晶分子が電場の向きにそろってねじれがなくなり，偏光方向が回転しないため光は遮られて暗く"
     "なる．画素ごとに電圧を制御し，カラーフィルタと組み合わせることで画像を表示する．")
para("(6)　白熱電球を用いた場合", sa=2, keep=True)
body("白熱電球の光はさまざまな波長を含み，位相もそろっていない．回折や干渉の位置（暗線・明線の位置）は波長に"
     "比例するので，波長ごとにしまの位置がずれて重なり，中央は白く，その外側に色のついた（内側が青，外側が赤の）"
     "ぼやけたしまが数本見えるだけになる．また光源の大きさが有限で光の向きもそろっていないため，しまのコントラ"
     "ストは低くなる．")
para("(7)　日常生活で見られる干渉や回折", sa=2, keep=True)
body("シャボン玉や水面に浮いた油膜が色づいて見えるのは，薄膜の表と裏で反射した光の干渉による．CDやDVDの記録面が"
     "虹色に見えるのは，細かく並んだ溝が回折格子のはたらきをするためである．ほかにも，街灯を目を細めて見たときの"
     "光の筋，カメラのレンズの反射防止膜，鳥の羽やモルフォチョウの羽の構造色などがある．")
para("(8)　ビームの広がり", sa=2, keep=True)
body(f"広がり角1.2 mradをビーム全体の広がり（全角）とすると，10 m進む間に直径は 10 m×1.2×10⁻³＝12 mm 増える．"
     f"したがって直径は 0.7 mm＋12 mm≈{beam10:.0f} mm になる．1.2 mrad を片側の角（半角）とみなした場合は，"
     f"0.7＋2×12≈25 mm となる．")

heading("B.7　まとめ")
body(f"レーザ光を用いて，単スリット・ピンホールによる回折と，複スリット・回折格子による干渉を観測した．スクリーン"
     f"上のしまの間隔と光の波長から，0.1 mm程度のスリットの幅やピンホールの直径を求めることができ，光が波として"
     f"ふるまうことを実感できた．回折格子では公称値との差が{abs(g_n - GRATING_NOM) / GRATING_NOM * 100:.1f} %と小さかった"
     f"一方，ピンホールでは差が大きく，ぼやけた暗環の位置をどこと判断するかが結果を大きく左右した．しまをいくつか"
     f"まとめて測る，暗い部分の中央を読むといった測り方の工夫が精度に直結することが分かった．")

C.references([
    "[1] 東京理科大学，“物理学実験 実験指導書 6．レーザ光の回折と干渉の実験，”2026．",
    "[2] E. Hecht, “Optics,” 5th ed., Pearson, 2017.",
])

if STANDALONE:
    out = os.path.join(HERE, "レーザ光の回折と干渉の実験_報告書.docx")
    C.save(out)
    print("saved:", out)
