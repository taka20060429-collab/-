"""図A-4（1/a と 1/b の関係）を、式入りの Excel ファイルとして作る。"""
import os
from openpyxl.chart import Reference, Series
from openpyxl.chart.marker import Marker
from openpyxl.styles import Font, PatternFill
import common as C
# lens.py と同じ値（凸レンズ・レンズの公式）
T_CONVEX = 0.495
CONVEX = [(22.370, 41.600, 62.620, 104.620), (22.555, 32.000, 53.300, 153.810),
          (22.400, 55.310, 76.705, 102.000), (22.430, 81.015, 102.340, 116.200)]

ws = C.new_sheet("図A-4")
ws.append(["回", "x1 [cm]", "x2 [cm]", "x3 [cm]", "x4 [cm]", "a [cm]", "b [cm]",
           "1/a [cm⁻¹]", "1/b [cm⁻¹]", "f [cm]"])
ws["L1"], ws["M1"] = "レンズの厚さ t [cm]", T_CONVEX
for i, (x1, x2, x3, x4) in enumerate(CONVEX, start=2):
    ws.append([i - 1, x1, x2, x3, x4,
               f"=ABS(D{i}-B{i})-$M$1/2", f"=ABS(E{i}-C{i})-$M$1/2",
               f"=1/F{i}", f"=1/G{i}", f"=1/(H{i}+I{i})"])
ws["I7"], ws["J7"] = "f の平均", "=AVERAGE(J2:J5)"
ws["L3"] = "公式の直線用（2点）"
ws["L4"], ws["M4"] = "1/a", "1/b"
ws["L5"], ws["M5"] = 0, "=1/$J$7"
ws["L6"], ws["M6"] = "=1/$J$7", 0
for c in ws[1]:
    c.font = Font(bold=True)
for addr in ("M1",):
    ws[addr].fill = PatternFill("solid", fgColor="FFF2CC")
for col, w in zip("ABCDEFGHIJKLM", [5, 9, 9, 9, 9, 9, 9, 11, 11, 9, 2, 18, 10]):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2, max_row=7, min_col=6, max_col=13):
    for c in row:
        c.number_format = "0.00000" if c.column_letter in "HILM" else "0.00"

ch = C.base_chart("1/a [cm⁻¹]", "1/b [cm⁻¹]")
s_l = Series(Reference(ws, min_col=13, min_row=5, max_row=6), Reference(ws, min_col=12, min_row=5, max_row=6),
             title="レンズの公式")
s_l.marker = Marker(symbol="none"); C.line(s_l, C.ORANGE)
s_p = Series(Reference(ws, min_col=9, min_row=2, max_row=5), Reference(ws, min_col=8, min_row=2, max_row=5),
             title="測定値")
s_p.marker = C.marker("circle", 7, C.BLUE); C.line(s_p, none=True)
ch.series.append(s_l); ch.series.append(s_p)
for ax in (ch.x_axis, ch.y_axis):
    ax.scaling.min, ax.scaling.max, ax.majorUnit = 0, 0.05, 0.005
xlsx, _, _ = C.native(ws, ch, "B10")
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "図A-4_1a_1bグラフ.xlsx")
open(out, "wb").write(xlsx)
print(out)
