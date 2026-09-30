# -*- coding: utf-8 -*-
"""実験レポート改善のまとめ（文書/Word）生成スクリプト。A4縦・黒一色・約4ページ。"""
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

IMGDIR = os.path.join(os.path.dirname(__file__), "img")
BODY_FONT = "ＭＳ 明朝"; HEAD_FONT = "ＭＳ ゴシック"; ASCII_FONT = "Times New Roman"

doc = Document()
st = doc.styles["Normal"]
st.font.name = ASCII_FONT; st.font.size = Pt(10.5); st.font.color.rgb = RGBColor(0, 0, 0)
st._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
st.paragraph_format.line_spacing = 1.15; st.paragraph_format.space_after = Pt(0)
sec = doc.sections[0]
sec.page_width = Mm(210); sec.page_height = Mm(297)
sec.top_margin = Mm(22); sec.bottom_margin = Mm(20); sec.left_margin = Mm(24); sec.right_margin = Mm(24)


def srf(run, ascii_font=ASCII_FONT, ea=BODY_FONT, size=None, bold=False):
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    rf.set(qn("w:ascii"), ascii_font); rf.set(qn("w:hAnsi"), ascii_font); rf.set(qn("w:eastAsia"), ea)
    if size:
        run.font.size = Pt(size)
    run.font.bold = bold; run.font.color.rgb = RGBColor(0, 0, 0)


def para(text="", size=10.5, bold=False, align=None, ea=BODY_FONT, ascii_f=ASCII_FONT,
         sb=0, sa=4, indent=None, line=1.3):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format; pf.space_before = Pt(sb); pf.space_after = Pt(sa); pf.line_spacing = line
    if indent is not None:
        pf.first_line_indent = Pt(indent)
    if text:
        r = p.add_run(text); srf(r, ascii_f, ea, size, bold)
    return p


def heading(num, text, level=1):
    sizes = {1: 13.5, 2: 11.5}; sb = {1: 12, 2: 8}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(sb[level]); p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{num}　{text}" if text else num); srf(r, HEAD_FONT, HEAD_FONT, sizes[level], True)
    return p


def body(text, indent=10.5):
    return para(text, size=10.5, sa=5, indent=indent, line=1.35)


def ba(before, after, aim):
    """改善前→改善後→ねらい の3行ブロック。"""
    for lbl, txt in [("改善前：", before), ("改善後：", after), ("ねらい：", aim)]:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Mm(4); p.paragraph_format.line_spacing = 1.3
        r1 = p.add_run(lbl); srf(r1, HEAD_FONT, HEAD_FONT, 10.5, True)
        r2 = p.add_run(txt); srf(r2, ASCII_FONT, BODY_FONT, 10.5, False)


def add_table(headers, rows, colw, fs=9.5):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers)); t.style = "Table Grid"; t.alignment = 1
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]; c.text = ""; rr = c.paragraphs[0].add_run(h)
        srf(rr, HEAD_FONT, BODY_FONT, fs, True); c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.rows[i + 1].cells[j]; c.text = ""; rr = c.paragraphs[0].add_run(v)
            srf(rr, ASCII_FONT, BODY_FONT, fs, False)
    for r in t.rows:
        for j, w in enumerate(colw):
            r.cells[j].width = Mm(w)
    return t


def figure(path, cap, w=160):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    if os.path.exists(path):
        run.add_picture(path, width=Mm(w))
    cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_after = Pt(8)
    cr = cp.add_run(cap); srf(cr, ASCII_FONT, BODY_FONT, 9.5, False)


def footer_pn():
    p = sec.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run = p.add_run()
    a = OxmlElement("w:fldChar"); a.set(qn("w:fldCharType"), "begin")
    b = OxmlElement("w:instrText"); b.set(qn("xml:space"), "preserve"); b.text = "PAGE"
    c = OxmlElement("w:fldChar"); c.set(qn("w:fldCharType"), "end")
    run._r.append(a); run._r.append(b); run._r.append(c); srf(run, ASCII_FONT, BODY_FONT, 10)


C = WD_ALIGN_PARAGRAPH.CENTER
# ===== タイトル =====
para("実験レポート改善のまとめ", size=19, bold=True, align=C, ea=HEAD_FONT, ascii_f=HEAD_FONT, sb=6, sa=6)
para("―「受動部品の種類と製作」レポートの制作過程における改善点 ―", size=12, align=C, sa=10)
para("作成者：4326091　森木崇允　／　作成日：2026年9月30日", size=10.5, align=C, sa=12)

# ===== 1 はじめに =====
heading("1", "はじめに")
body("本書は、「受動部品の種類と製作」実験レポートを作成するにあたり、初回の試作から完成版に至るまで"
     "に行った主な改善点を、改善前・改善後の対比とともにまとめたものである。あわせて、これらの学びを"
     "後続の「ネットワーク実習」レポートにどのように反映したかを記す。試作段階のレポートと完成版とでは、"
     "出力形式・構成・分量・内容の具体性の各面で大きな違いが生じた。")

# ===== 2 改善点の一覧 =====
heading("2", "改善点の一覧")
add_table(
    ["#", "項目", "改善前", "改善後"],
    [["1", "出力形式", "PowerPoint（スライド）", "Word（A4縦の文書）"],
     ["2", "配色", "4色のカラフルな配色", "黒一色"],
     ["3", "全体構成", "独自の15スライド構成", "過去レポート準拠（目的／原理／方法／結果／考察／まとめ／参考文献）"],
     ["4", "写真", "なし。後から貼った際も切り取りミスで机だけ", "実験写真を本体が写るよう切り直して貼付"],
     ["5", "実験データ", "ダミー値・【要記入】の空欄", "自分の実測値を反映"],
     ["6", "考察", "教科書的な一般論のみ", "理論値計算＋自分の言葉の所感"],
     ["7", "原理の分量", "薄い（全5ページ）", "分類表・カラーコード表を追加し全7〜8ページ"],
     ["8", "表紙情報", "氏名・学籍番号が空欄", "記入済み"]],
    colw=[8, 24, 60, 68])
para("", sa=2)
figure(os.path.join(IMGDIR, "comp_format.jpg"), "図1　出力形式・配色の改善（左：改善前／右：改善後）", w=150)

# ===== 3 各改善の詳細 =====
heading("3", "各改善の詳細（改善前 → 改善後）")

heading("3-1", "出力形式：PowerPoint → Word文書", level=2)
ba("最初はスライド形式のPowerPointで作成した。",
   "過去レポートが文書形式であったため、A4縦のWord文書に作り直した。",
   "提出形式・体裁を過去レポートに合わせ、採点者が見慣れた形にするため。")

heading("3-2", "配色：カラフル → 黒一色", level=2)
ba("紺・赤・青緑・緑の4色で章を色分けしていた。",
   "本文・見出しとも黒一色に統一した。",
   "学術的な実験レポートとしての体裁を整え、印刷・提出に適した落ち着いた見た目にするため。")

heading("3-3", "全体構成：独自構成 → 過去レポート準拠", level=2)
ba("独自の15スライド構成だった。",
   "「1 目的／2 原理／3 方法／4 結果／5 考察／6 まとめ／参考文献」という過去レポートと同じ構成に統一した。",
   "評価基準に沿った標準的なレポート構成にするため。")

heading("3-4", "写真：なし・切り取りミス → 適切に貼付", level=2)
ba("当初は写真がなく、後から貼った際もコンデンサ写真の切り取り位置を誤り、机だけが写っていた。",
   "鉛筆抵抗カード・マンガニン線抵抗器・コンデンサの各写真を、本体がはっきり写るよう切り直して結果の該当箇所に貼付した。",
   "自作した部品を実物写真で示し、結果の説得力と実験の実在性を高めるため。")

heading("3-5", "実験データ：ダミー・空欄 → 実測値", level=2)
ba("数値がダミーや【要記入】のままだった。",
   "鉛筆抵抗（①2100Ω／②20kΩ／③1000kΩ）、マンガニン線抵抗（17Ω）、コンデンサ（23nF）などの実測値を反映し、インダクタは未実施として明記した。",
   "自分の実験結果に基づく正確なレポートにするため。")

heading("3-6", "考察：一般論 → 計算の具体化＋所感", level=2)
ba("考察が教科書的な一般論にとどまっていた。",
   "R＝ρL／S から鉛筆芯の筆跡の厚さ（約1.9×10⁻⁵ m）を算出し、C＝εS／d の理論値（約2.7nF）と実測23nFを比較するなど計算を具体化した。さらに「マンガニン線抵抗が想像より低くて驚いた」等の自分の言葉の所感を加えた。",
   "考察の説得力を高め、機械的でない“自分のレポート”らしさを出すため。")

heading("3-7", "原理の分量：薄い → 大幅増量", level=2)
ba("原理が薄く、全体で5ページ程度で、過去レポートの原理との分量差が大きかった。",
   "指導書・講義資料をもとに抵抗器・コンデンサ・インダクタの各種類を詳述し、「表1 抵抗器の分類」「表2 カラーコード表」を追加して全体で7〜8ページに増量した。",
   "過去レポートと同等の充実度にし、理解の深さを示すため。")

heading("3-8", "表紙情報：空欄 → 記入", level=2)
ba("氏名・学籍番号が【要記入】のままだった。",
   "学籍番号4326091・氏名森木崇允・実験日・提出日を記入した。",
   "提出に必要な情報を整えるため。")

# ===== 4 改善の進め方 =====
heading("4", "改善の進め方")
body("本レポートは一度で完成させたのではなく、反復的な手順で仕上げた。まず過去レポートを基準として"
     "全体の構成と体裁を組み立て、次に実験写真や実測データを順次反映し、さらに指摘を受けた点（写真の"
     "切り取りミス、原理の分量不足など）を一つずつ修正していった。基準となる過去レポートと自分の実験"
     "結果を突き合わせながら改善を重ねたことで、抜けや体裁の崩れを段階的に解消でき、最終的に完成度の"
     "高いレポートにまとめることができた。")

# ===== 5 ネットワーク実習への反映 =====
heading("5", "ネットワーク実習レポートへの反映")
body("以上の学びを踏まえ、後続の「ネットワーク実習」レポートは最初から、黒一色のWord文書・過去レポート"
     "準拠の構成（同レポートでは目的→考察→まとめ）・結果写真の貼付・自分の言葉の所感、という方針で作成"
     "した。これにより試行錯誤の必要が減り、短時間で完成度の高いレポートを作成できた。改善で得た方針が、"
     "別テーマのレポートにもそのまま活かせることが確認できた。")

# ===== 6 まとめ =====
heading("6", "まとめ")
body("一連の改善を通じて、実験レポートの作成では次の4点が完成度を大きく左右すると分かった。"
     "①提出形式・体裁を過去レポートに合わせること、②自分の実測データと写真を正確に反映すること、"
     "③考察を計算と自分の言葉で具体化すること、④原理などの分量を十分に確保すること、である。"
     "今後のレポート作成においても、この方針を踏襲していきたい。")

footer_pn()
out = os.path.join(os.path.dirname(__file__), "実験レポート改善のまとめ.docx")
doc.save(out); print("saved:", out)
