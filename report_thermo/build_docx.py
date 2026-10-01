# -*- coding: utf-8 -*-
"""熱電対の検定実験 報告書（Word）生成スクリプト。
使い方: python3 build_docx.py [データファイル]   （省略時は同じフォルダの data.py）
測定データがそろっていない項目は【要記入】として出力する。"""
import hashlib
import importlib.util
import math
import os
import re
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

import analysis as A

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "img")
EQ = os.path.join(IMG, "eq")
os.makedirs(EQ, exist_ok=True)

data_path = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "data.py")
_spec = importlib.util.spec_from_file_location("data", data_path)
D = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(D)
out_path = sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, "熱電対の検定実験_報告書.docx")

COMPLETE = A.is_complete(D)
R = A.compute(D) if COMPLETE else None
A.make_static_figs(IMG)
if COMPLETE:
    A.make_data_figs(R, D, IMG)

BODY, HEAD, LATIN = "ＭＳ 明朝", "ＭＳ ゴシック", "Times New Roman"
TODO = "【要記入】"

# ================================================================ 文書の基本設定
doc = Document()
st = doc.styles["Normal"]
st.font.name = LATIN
st.font.size = Pt(10.5)
st.font.color.rgb = RGBColor(0, 0, 0)
st._element.rPr.rFonts.set(qn("w:eastAsia"), BODY)
st.paragraph_format.line_spacing = 1.2
st.paragraph_format.space_after = Pt(0)
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin = sec.bottom_margin = Mm(25)
sec.left_margin = sec.right_margin = Mm(25)
TEXT_W = 160  # mm


def font(run, latin=LATIN, ea=BODY, size=None, bold=False, italic=False):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:ascii"), latin)
    rf.set(qn("w:hAnsi"), latin)
    rf.set(qn("w:eastAsia"), ea)
    if size:
        run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def para(text="", size=10.5, bold=False, align=None, ea=BODY, latin=LATIN, sb=0, sa=4,
         indent=None, left=None, line=1.35, keep=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(sb), Pt(sa), line
    if indent is not None:
        pf.first_line_indent = Pt(indent)
    if left is not None:
        pf.left_indent = Mm(left)
    if keep:
        pf.keep_with_next = True
    if text:
        add_text(p, text, latin, ea, size, bold)
    return p


_MK = re.compile(r"(_\{[^}]*\}|\^\{[^}]*\})")


def add_text(p, text, latin=LATIN, ea=BODY, size=None, bold=False):
    """_{..} を下付き、^{..} を上付きの文字として書き込む。"""
    for tok in _MK.split(text):
        if not tok:
            continue
        if tok[:2] in ("_{", "^{"):
            r = p.add_run(tok[2:-1])
            font(r, latin, ea, size, bold)
            if tok[0] == "_":
                r.font.subscript = True
            else:
                r.font.superscript = True
        else:
            font(p.add_run(tok), latin, ea, size, bold)


def body(text):
    return para(text, indent=10.5)


def item(text, left=4):
    p = para(text, left=left, line=1.35, sa=3)
    p.paragraph_format.first_line_indent = Mm(-4)
    p.paragraph_format.left_indent = Mm(left + 4)
    return p


def heading(text, level=1):
    size = {1: 13, 2: 11.5, 3: 10.5}[level]
    p = para(sb={1: 14, 2: 9, 3: 6}[level], sa=5, keep=True)
    font(p.add_run(text), HEAD, HEAD, size, True)
    return p


def todo(text):
    p = para(left=4, sa=4)
    font(p.add_run(text), LATIN, BODY, 10.5, True)
    return p


# ---------------------------------------------------------------- 図・表・式の番号
CNT = {"fig": 0, "tab": 0, "eq": 0}


def nxt(kind):
    CNT[kind] += 1
    return CNT[kind]


def figure(path, caption, num, width=130):
    p = para(align=WD_ALIGN_PARAGRAPH.CENTER, sb=6, sa=2, keep=True)
    p.add_run().add_picture(path, width=Mm(width))
    c = para(align=WD_ALIGN_PARAGRAPH.CENTER, sa=10)
    font(c.add_run(f"図{num}　{caption}"), LATIN, BODY, 9.5)


def table(caption, num, headers, rows, widths, size=9.5, align_cols=None):
    c = para(align=WD_ALIGN_PARAGRAPH.CENTER, sb=6, sa=3, keep=True)
    font(c.add_run(f"表{num}　{caption}"), LATIN, BODY, 9.5)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ""
        q = cell.paragraphs[0]
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(q.add_run(h), LATIN, BODY, size, False)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = t.rows[i + 1].cells[j]
            cell.text = ""
            q = cell.paragraphs[0]
            q.alignment = (align_cols[j] if align_cols else WD_ALIGN_PARAGRAPH.CENTER)
            font(q.add_run(str(v)), LATIN, BODY, size, False)
    t.autofit = False
    for j, w in enumerate(widths):
        t.columns[j].width = Mm(w)
    for r in t.rows:
        for j, w in enumerate(widths):
            r.cells[j].width = Mm(w)
    para(sa=6)
    return t


def _render_tex(tex, size=12.5):
    name = "eq" + hashlib.md5(f"{tex}|{size}".encode()).hexdigest()[:12] + ".png"
    path = os.path.join(EQ, name)
    if not os.path.exists(path):
        fig = plt.figure(figsize=(0.01, 0.01))
        fig.text(0, 0, f"${tex}$", fontsize=size)
        fig.savefig(path, dpi=600, bbox_inches="tight", pad_inches=0.03, transparent=True)
        plt.close(fig)
    from PIL import Image
    w, h = Image.open(path).size
    return path, w / 600 * 25.4  # 自然な大きさ [mm]


def eq(tex, numbered=True, size=12.5):
    """式を中央に、式番号を右端に置く。"""
    path, wmm = _render_tex(tex, size)
    wmm = min(wmm, TEXT_W - 22)
    p = para(sb=3, sa=5, line=1.0)
    ts = p.paragraph_format.tab_stops
    ts.add_tab_stop(Mm(TEXT_W / 2), WD_TAB_ALIGNMENT.CENTER)
    ts.add_tab_stop(Mm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
    font(p.add_run("\t"))
    p.add_run().add_picture(path, width=Mm(wmm))
    n = None
    if numbered:
        n = nxt("eq")
        font(p.add_run(f"\t({n})"))
    return n


# ---------------------------------------------------------------- 数値の書式
SUP = str.maketrans("0123456789-", "⁰¹²³⁴⁵⁶⁷⁸⁹⁻")


def sci(x, sig=4):
    """本文用: 4.100×10⁻² のような表記。"""
    if x == 0:
        return "0"
    e = int(math.floor(math.log10(abs(x))))
    if -2 <= e <= 3:
        d = max(sig - 1 - e, 0)
        return f"{x:.{d}f}"
    m = x / 10 ** e
    return f"{m:.{sig - 1}f}×10{str(e).translate(SUP)}"


def tsci(x, sig=4):
    """式画像用（mathtext）。"""
    s = sci(x, sig)
    if "×" in s:
        m, e = s.split("×10")
        e = e.translate(str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹⁻", "0123456789-"))
        return f"{m}\\times 10^{{{e}}}"
    return s


def v(x, fmt):
    return TODO if x is None else format(x, fmt)


# ================================================================ 表紙
M = D.META
for _ in range(4):
    para(sa=4)
para(f"第 {M['kai'] if M['kai'] else '　'} 回　物理学実験　報告書", size=16, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, ea=HEAD, latin=HEAD, sa=26)
para("実験題目：熱電対の検定実験", size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
     ea=HEAD, latin=HEAD, sa=40)
cover = [
    f"実験年月日：{M['date'] or TODO}",
    f"報告書提出日：{M['submit'] or TODO}",
    f"班名：{M['han'] or TODO}",
    f"実験者氏名：{M['name']}　　学年：{M['grade']}　　学籍番号：{M['student_id']}",
]
for s in cover:
    para(s, size=11.5, left=28, sa=10)
para("共同実験者：", size=11.5, left=28, sa=6)
if M["coworkers"]:
    for nm, sid in M["coworkers"]:
        para(f"氏名：{nm}　　学籍番号：{sid}", size=11, left=38, sa=4)
else:
    para(f"氏名：{TODO}　　学籍番号：{TODO}", size=11, left=38, sa=4)
doc.add_page_break()

# ================================================================ 1 目的
heading("1　目的")
body("熱電対を用いた温度測定の原理と方法を理解したうえで，融点が既知である3種類の金属（Sn，Pb，Zn）"
     "の凝固点，ならびに0 ℃と室温を定点として，クロメル・アルメル熱電対の熱起電力と温度の関係を"
     "校正（検定）する．あわせて，測定値から関係式を決める手法である最小二乗法を理解し，1次式と2次式"
     "による近似の結果を比較・検討する．")

# ================================================================ 2 原理
heading("2　原理")
heading("2.1　熱電対とゼーベック効果", 2)
body("2種類の異なる金属線A，Bの両端を接合して閉回路をつくり，二つの接合点の温度を異なる値に保つと，"
     "回路に起電力が生じて電流が流れる．この現象をゼーベック効果といい，生じる起電力を熱起電力という[1]．"
     "熱起電力の向きと大きさは，金属の組み合わせと二つの接合点の温度だけで決まり，線の長さや太さ，"
     "途中の温度分布には左右されない．また，同じ均質な金属だけで回路をつくっても熱起電力は生じない"
     "（均質回路の法則）[1]．")
body("この現象は，金属中の自由電子の振る舞いから定性的に説明できる．金属棒の一端を高温にすると，"
     "高温側の自由電子は大きな運動エネルギーをもつため低温側へ拡散し，低温側に電子がたまって棒の内部に"
     "電界が生じる．電子1個あたりの比熱をCₑ，電気素量をeとすると，温度差ΔTによる運動エネルギーの差"
     "CₑΔTと電位差ΔVの間には eΔV＝CₑΔT の関係が成り立つと考えられ，単位温度差あたりの電位差 "
     "S＝ΔV/ΔT＝Cₑ/e をゼーベック係数（熱電能）という[1]．Cₑは金属ごとに異なるため，異なる2種類の"
     "金属を接合すると両者の電位差が打ち消し合わず，その差が熱起電力として現れる．金属A，Bの"
     "ゼーベック係数をそれぞれS_{A}，S_{B}，測温接点の温度をT，基準接点の温度をT₀とすると，熱起電力Eは")
n_seebeck = eq(r"E=\int_{T_0}^{T}(S_A-S_B)\,dT\ \approx\ S_{AB}\,(T-T_0)")
body(f"と表される．ここで S_{{AB}}＝S_{{A}}－S_{{B}} は熱電対としてのゼーベック係数である．S_{{AB}} が温度によらず"
     f"一定であればEは温度差に比例するが，実際には S_{{AB}} は温度によってわずかに変化するため，"
     f"Eと温度の関係は厳密には直線にならない．なお，ゼーベック効果を厳密に扱うにはトムソン効果や"
     f"ペルチェ効果も考慮する必要がある[1]．")

heading("2.2　熱電対による温度測定", 2)
f_circ = nxt("fig")
body("熱起電力は二つの接合点の温度差と一対一に対応するので，一方の接合点（基準接点）を既知の温度に"
     "保てば，熱起電力からもう一方の接合点（測温接点）の温度を求めることができる．基準接点を室温に"
     "置くと室温の変動がそのまま誤差になるため，本実験では氷と水が共存して0 ℃に保たれる氷水中に"
     "基準接点を置く．基準接点から電圧計までは銅線で配線するが，回路に第3の金属を挿入してもその両端の"
     "温度が等しければ熱起電力は変わらないという中間金属の法則により，測定される起電力は熱電対の金属の"
     f"種類と温度T，T₀だけで決まる[1]．測定の構成を図{f_circ}に示す．基準接点を0 ℃としたときの温度と"
     "熱起電力の対応は熱起電力表として規格化されており[2]，本実験で用いるクロメル（＋脚）・アルメル"
     "（－脚）熱電対はK熱電対と呼ばれる．")
figure(os.path.join(IMG, "fig_circuit.png"), "熱電対による温度測定の構成（模式図）", f_circ, 125)

heading("2.3　金属の凝固点を利用した熱電対の検定", 2)
f_cs = nxt("fig")
body("熱電対は素線の不均質や劣化などによって規格の熱起電力表からずれることがあるため，温度が正確に"
     "分かっている点（定点）で熱起電力を測定し，起電力と温度の関係を校正する必要がある．本実験では"
     "これを検定と呼ぶ[1]．定点としては純金属の凝固点（融点）を用いる．")
body("融解させた金属をゆっくり冷却すると，凝固が始まってから終わるまでの間は凝固に伴う潜熱が放出され"
     "続けるため，固相と液相が共存する間は温度がほぼ一定に保たれる．したがって，温度（熱起電力）の"
     f"時間変化を表す冷却曲線には，図{f_cs}の実線のような平坦な部分（以下，プラトー）が現れ，その温度が"
     "金属の融点に等しい．本実験では，Sn（231.97 ℃），Pb（327.52 ℃），Zn（419.58 ℃）の3種類の"
     "金属を用いる[1]．")
body("ただし，液体を静かに冷却すると，融点を下回っても凝固が始まらない過冷却が起こることがある．"
     "原子が規則正しく並んだ結晶ができるには，そのきっかけとなる結晶核の形成に必要なエネルギー障壁"
     "（活性化エネルギー）を越えなければならないためである[1]．過冷却の状態から凝固が始まると，放出"
     "される潜熱によって温度は融点まで急に戻り（再輝），その後プラトーが現れる（図"
     f"{f_cs}の破線）．そのため融点に対応する起電力は，過冷却で下がった部分ではなく，再上昇した後の"
     "プラトーから決める．なお，振動を与えたり不純物が存在したりすると結晶核ができやすくなるため，"
     "過冷却は起こりにくくなる．")
figure(os.path.join(IMG, "fig_cool_schematic.png"), "金属の冷却曲線（模式図）", f_cs, 120)

heading("2.4　最小二乗法による関係式の決定", 2)
body("検定では，測定した温度tと熱起電力Eの組から両者の関係式を求める．K熱電対の基準熱起電力は，"
     "温度の高次の多項式に指数関数の項を加えた式で表されるが[1][3]，本実験では0〜420 ℃程度の範囲に"
     "限って，次の1次式および2次式で近似する．")
n_lin = eq(r"E=at+b")
n_quad = eq(r"E=at^2+bt+c")
body("係数は最小二乗法で決める．N個の測定値の組 (x_{i}, y_{i})（i＝1, …, N）に対し，仮定した関係式"
     "f(x)による推定値を ŷ_{i}＝f(x_{i}) とすると，最小二乗法は残差 y_{i}－ŷ_{i} の二乗和")
n_Q = eq(r"Q=\sum_{i=1}^{N}\left(y_i-\hat{y}_i\right)^2")
body("を最小にする係数を求める方法である．本実験ではxが温度t，yが熱起電力Eに対応する．なお，"
     "ゼーベック係数Sとの混同を避けるため，本報告書では残差平方和をQで表す．Qは係数について下に凸な"
     "2次関数であるから，Qを各係数で偏微分して0とおいた連立方程式（正規方程式）の解が最小値を与える．")
para("(i)　1次式 f(x)＝ax＋b の場合", left=0, sa=2, keep=True)
n_ne1 = eq(r"\dfrac{\partial Q}{\partial a}=-2\sum_{i=1}^{N}x_i\,(y_i-ax_i-b)=0,\quad"
           r"\dfrac{\partial Q}{\partial b}=-2\sum_{i=1}^{N}(y_i-ax_i-b)=0")
body(f"式({n_ne1})の第2式の両辺を－2Nで割ると ȳ－ax̄－b＝0 となり b＝ȳ－ax̄ が得られる．これを第1式に"
     "代入して整理すると，係数は")
n_ab = eq(r"a=\dfrac{S_{xy}}{S_x},\qquad b=\bar{y}-a\bar{x}")
body("となる．ここで x̄，ȳ はそれぞれx，yの平均であり，S_{x} はxの分散，S_{xy} はxとyの共分散で，")
n_S1 = eq(r"S_x=\overline{x^2}-\bar{x}^{\,2},\qquad S_{xy}=\overline{xy}-\bar{x}\,\bar{y}")
body("である．上線はN個の値の平均を表す．")
para("(ii)　2次式 f(x)＝ax²＋bx＋c の場合", left=0, sa=2, keep=True)
body("同様にQをa，b，cで偏微分して0とおく．cについての式から c＝ȳ－a·(x²の平均)－bx̄ が得られ，"
     "これを残りの2式に代入すると")
n_ne2 = eq(r"aS_{x^2}+bS_{xx^2}=S_{x^2y},\qquad aS_{xx^2}+bS_x=S_{xy}")
body("となる．これをa，bについて解くと，")
eq(r"a=\dfrac{S_xS_{x^2y}-S_{xy}S_{xx^2}}{S_xS_{x^2}-S_{xx^2}^{\,2}}", numbered=False)
eq(r"b=\dfrac{S_{xy}S_{x^2}-S_{xx^2}S_{x^2y}}{S_xS_{x^2}-S_{xx^2}^{\,2}}", numbered=False)
n_abc = eq(r"c=\bar{y}-a\overline{x^2}-b\bar{x}")
body("が得られる．ここで")
n_S2 = eq(r"S_{x^2}=\overline{x^4}-\left(\overline{x^2}\right)^2,\quad"
          r"S_{x^2y}=\overline{x^2y}-\overline{x^2}\,\bar{y},\quad"
          r"S_{xx^2}=\overline{x^3}-\bar{x}\,\overline{x^2}", size=12)
body(f"であり，それぞれx²の分散，x²とyの共分散，xとx²の共分散を表す[1]．")

# ================================================================ 3 実験方法
heading("3　実験方法")
f_setup = nxt("fig")
heading("3.1　準備", 2)
body(f"図{f_setup}のように装置を結線した．誤った結線のまま電気炉に通電すると危険であるため，担当者に"
     "結線を確認してもらうまでスライダックの電源プラグは差し込まなかった．")
item("(1)　デュワービンに氷と水道水を入れて氷水とし，基準接点を0 ℃に保った．")
item("(2)　電気炉のヒータの導線を，交流電流計（10 Aレンジ）を介してスライダックの出力端子に接続した．")
item("(3)　熱電対の端をデュワービンのTHERMOCOUPLE端子に接続した．金属光沢の強いクロメルを＋側，"
     "アルメルを－側とし，極性を誤らないようにした．")
item("(4)　デジタルボルトメータのV端子とCOM端子をデュワービンのMETER端子に接続した．")
figure(os.path.join(IMG, "fig_setup.png"), "実験装置の構成", f_setup, 135)
heading("3.2　実験1：室温での熱起電力の測定", 2)
body("室温を室内の温度計で読み取った．次に，熱電対を加熱していない状態でデジタルボルトメータの表示が"
     "安定するのを待ち，その値を室温での熱起電力として記録した．")
heading("3.3　実験2：0 ℃での熱起電力の測定", 2)
body("予備の保護管に熱電対を入れ，デュワービンのふたの穴から氷水中に差し込んだ（熱電対がぬれない"
     "ようにした）．表示が安定したときの値を0 ℃での熱起電力として記録した．理想的には0 mVとなる．")
heading("3.4　実験3：金属の融点に対応する熱起電力の測定", 2)
body("Sn，Pb，Znのそれぞれについて，次の手順で測定した．")
item("(1)　熱電対を保護管に入れ，先端が保護管の内側の底まで届いていることを確かめてから，試料の入った"
     "タンマン管に差し込んだ．先端が試料の位置に届いていないと，試料ではなく周囲の空気の温度を測る"
     "ことになるためである．")
item("(2)　タンマン管を電気炉内の台座に傾かないように置き，熱電対の2本の素線どうしが接触していない"
     "ことを確認した．素線が途中で接触すると，その点が新たな測温接点になってしまうためである．")
item("(3)　スライダックのつまみが0 Vであることを確認してから電源を入れ，つまみを少しずつ回して電気炉に"
     "9.5 Aの電流を流して加熱した．加熱を始めた時点から15 s間隔で熱起電力を記録し続け，加熱中に上昇が鈍る"
     "点を融点の目安とした．")
item("(4)　融点を越えてから熱起電力がさらに1.0〜1.5 mV程度上昇したところでスライダックを0 Vに戻し，"
     "電流が0 Aになったことを確認した．")
item("(5)　余熱によってしばらく上昇したのち熱起電力は下がり始めた．その後も15 s間隔で記録を続け"
     "（冷却実験），プラトーを過ぎて熱起電力が融点に相当する値より1.0〜1.5 mV程度低下したところで"
     "測定を終えた．過冷却が起きた場合は，ピンセットでタンマン管に軽く振動を与えた．記録は加熱開始から"
     "通して行ったため，整理の際は熱起電力が最大となった時点を冷却の0 sとし，それ以降を冷却時の"
     "データとして扱った．")
body("なお，加熱中の電気炉，タンマン管および保護管は高温になるため，触れる必要がある場合は軍手と"
     "ピンセットを用いた．")

# ================================================================ 4 使用器具
heading("4　使用器具")
t_eq = nxt("tab")
body(f"使用した器具を表{t_eq}に示す．")
table("使用器具", t_eq, ["器具", "数量", "定格・備考"],
      [["クロメル・アルメル熱電対（K熱電対）", "1本", "＋脚：クロメル，－脚：アルメル"],
       ["標準金属試料（Sn，Pb，Zn）", "各1", "タンマン管に封入"],
       ["タンマン管・保護管", "各3本", ""],
       ["電気炉", "1台", ""],
       ["デュワービン", "1台", "基準接点（氷水，0 ℃）"],
       ["デジタルボルトメータ", "1台", "熱起電力の測定"],
       ["スライダック", "1台", "130 V，10 A"],
       ["交流電流計", "1台", "10 A"],
       ["ストップウォッチ，スタンド", "各1", ""],
       ["軍手，ピンセット", "一式", ""]],
      [62, 18, 60], align_cols=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
                                WD_ALIGN_PARAGRAPH.LEFT])

# ================================================================ 5 実験結果
heading("5　実験結果")
REFS_CITED = set()


def draft_results():
    heading("5.1　室温での熱起電力", 2)
    body(f"室温 t＝{TODO} ℃ において，熱起電力は E＝{TODO} mV であった．")
    heading("5.2　0 ℃での熱起電力", 2)
    body(f"0 ℃での熱起電力は E₀＝{TODO} mV であった．")
    heading("5.3　冷却曲線と融点に対応する熱起電力", 2)
    todo("【測定データ受領後に自動作成】Sn・Pb・Znの冷却時の熱起電力の表（15 s間隔），冷却曲線（横軸：時間[s]，"
         "縦軸：熱起電力[mV]），プラトー区間の判定（隣り合う測定値の差が0.05 mV以下の区間）と，その区間の"
         "平均・標準偏差による融点の起電力の決定．")
    heading("5.4　熱電対の検定（最小二乗法）", 2)
    todo("【測定データ受領後に自動作成】5つの定点（0 ℃，室温，Sn，Pb，Znの融点）の表，最小二乗法の計算表"
         "（t, E, t², tE, t³, t⁴, t²E とその平均），式(6)〜(10)に数値を代入した計算過程，近似直線・近似曲線の式，"
         "5点と近似直線・近似曲線を1つにまとめたグラフ，各点の残差．")


def cooling_table(m, num):
    c = R["cool"][m]
    pairs = [(f"{t:.0f}", f"{e:.3f}") for t, e in zip(c["t"], c["E"])]
    ncol = 3
    nrow = math.ceil(len(pairs) / ncol)
    rows = []
    for r in range(nrow):
        row = []
        for k in range(ncol):
            i = r + k * nrow
            row += list(pairs[i]) if i < len(pairs) else ["", ""]
        rows.append(row)
    table(f"{A.NAME_JP[m]}（{m}）の冷却時の熱起電力", num,
          ["時間 [s]", "熱起電力 [mV]"] * ncol, rows, [22, 31] * ncol, size=8.5)


def data_results():
    heading("5.1　室温での熱起電力", 2)
    body(f"室温 t＝{D.ROOM_TEMP:.1f} ℃ において，熱起電力は E＝{D.ROOM_EMF:.3f} mV であった．")
    heading("5.2　0 ℃での熱起電力", 2)
    z = D.ZERO_EMF
    body(f"0 ℃での熱起電力は E₀＝{z:.3f} mV であった．理想値0 mVとの差は{abs(z) * 1000:.0f} µVであり，"
         f"0 ℃付近のゼーベック係数（約{A.seebeck_std(0) * 1000:.1f} µV/℃[3]）を用いて温度に換算すると"
         f"約{abs(z) / A.seebeck_std(0):.1f} ℃に相当する．")

    heading("5.3　冷却曲線と融点に対応する熱起電力", 2)
    nums = {m: (nxt("tab"), nxt("fig")) for m in A.ORDER}
    tabs = "，".join(f"表{nums[m][0]}" for m in A.ORDER)
    figs = "，".join(f"図{nums[m][1]}" for m in A.ORDER)
    pk = "，".join(f"{A.NAME_JP[m]}では{R['cool'][m]['t_peak']:.0f} s後に{R['cool'][m]['E_peak']:.3f} mV"
                  for m in A.ORDER if R["cool"][m]["t_peak"] is not None)
    if pk:
        body(f"熱起電力は加熱開始から，{pk}で最大となった．3.4節で述べたとおり，この時点を冷却の0 sとし，"
             f"以降のデータを冷却時の結果として整理した．")
    body(f"Sn，Pb，Znの冷却時に15 s間隔で記録した熱起電力を{tabs}に，冷却曲線を{figs}に示す．冷却曲線の"
         f"縦軸は，プラトー付近の変化が読み取れるよう測定値の範囲に合わせて拡大した．融点に対応する起電力"
         f"は次のように決めた．隣り合う測定値（15 s間隔）の差が{A.PLATEAU_THR:.2f} mV（K熱電対で約"
         f"{A.PLATEAU_THR / A.seebeck_std(300):.1f} ℃に相当）以下の状態が最も長く続いた区間をプラトーとみなし，"
         f"その区間の測定値の平均を融点に対応する熱起電力とした．また，その区間の標準偏差を，プラトーの"
         f"平坦さ（値のばらつき）の目安として併記した．")
    summary = []
    for m in A.ORDER:
        c = R["cool"][m]
        tn, fn = nums[m]
        heading(f"5.3.{A.ORDER.index(m) + 1}　{A.NAME_JP[m]}（{m}，融点 {A.MELT[m]:.2f} ℃）", 3)
        cooling_table(m, tn)
        figure(os.path.join(IMG, f"fig_cool_{m}.png"),
               f"{A.NAME_JP[m]}（{m}）の冷却曲線（灰色の帯：融点とみなした区間）", fn, 130)
        vals = " ＋ ".join(f"{x:.3f}" for x in c["values"])
        body(f"図{fn}より，{c['t0']:.0f}〜{c['t1']:.0f} s の{c['n']}点がプラトーにあたる．この区間の平均は")
        para(f"（{vals}）／{c['n']} ＝ {c['mean']:.3f} mV", left=6, sa=4, size=10)
        txt = (f"であり，標準偏差は{c['sd']:.3f} mVであった．よって，{A.NAME_JP[m]}の融点 "
               f"{A.MELT[m]:.2f} ℃ に対応する熱起電力を {c['mean']:.3f} mV とした．")
        sc = c["supercool"]
        if sc:
            txt += (f"なお，{sc['t_min'] * D.DT:.0f} s付近で熱起電力が{sc['E_min']:.3f} mVまで下がったのち，"
                    f"{sc['rise']:.3f} mV上昇してプラトーに移っており，過冷却が起きていたことが分かる．"
                    f"過冷却で下がった部分は融点の決定には用いなかった．")
        body(txt)
        summary.append([f"{A.NAME_JP[m]}（{m}）", f"{A.MELT[m]:.2f}", f"{c['t0']:.0f}〜{c['t1']:.0f}",
                        f"{c['n']}", f"{c['mean']:.3f}", f"{c['sd']:.3f}"])
    ts = nxt("tab")
    body(f"以上の結果を表{ts}にまとめる．")
    table("融点に対応する熱起電力", ts,
          ["金属", "融点 [℃]", "プラトー区間 [s]", "点数", "熱起電力 [mV]", "標準偏差 [mV]"],
          summary, [26, 22, 30, 14, 28, 28], size=9)

    # ------------------------------------------------------------ 5.4 検定
    heading("5.4　熱電対の検定（最小二乗法）", 2)
    L = R["lsq"]
    t5 = nxt("tab")
    body(f"5.1〜5.3で得た5組の温度と熱起電力を表{t5}に示す．これらに対して，2.4節の最小二乗法により"
         f"近似直線（式({n_lin})）と近似曲線（式({n_quad})）を求めた．以下，x，yをそれぞれ温度t，熱起電力Eに"
         f"読み替え，N＝5 とする．")
    table("検定に用いた温度と熱起電力", t5, ["定点", "温度 t [℃]", "熱起電力 E [mV]"],
          [[lab, f"{t:.2f}" if t not in (0.0,) else "0", f"{e:.3f}"] for lab, t, e in R["points"]],
          [50, 35, 40])
    t6 = nxt("tab")
    body(f"計算に必要な各量とその平均を表{t6}に示す．")
    rows = []
    for i, r in enumerate(L["rows"], 1):
        rows.append([str(i), f"{r['x']:.2f}", f"{r['y']:.3f}", sci(r["x2"], 6), sci(r["xy"], 6),
                     sci(r["x3"], 6), sci(r["x4"], 6), sci(r["x2y"], 6)])
    rows.append(["平均", sci(L["xb"], 6), sci(L["yb"], 6), sci(L["x2b"], 6), sci(L["xyb"], 6),
                 sci(L["x3b"], 6), sci(L["x4b"], 6), sci(L["x2yb"], 6)])
    table("最小二乗法の計算に用いた値（t [℃]，E [mV]）", t6,
          ["i", "t", "E", "t²", "tE", "t³", "t⁴", "t²E"], rows, [10, 17, 15, 21, 19, 24, 27, 27], size=8)

    para("(i)　近似直線（1次式）", sa=2, keep=True)
    body(f"式({n_S1})より，")
    eq(rf"S_t=\overline{{t^2}}-\bar{{t}}^{{\,2}}={tsci(L['x2b'], 7)}-({tsci(L['xb'], 6)})^2={tsci(L['Sx'], 7)}",
       numbered=False, size=11)
    eq(rf"S_{{tE}}=\overline{{tE}}-\bar{{t}}\,\bar{{E}}={tsci(L['xyb'], 7)}-{tsci(L['xb'], 6)}\times{tsci(L['yb'], 6)}"
       rf"={tsci(L['Sxy'], 7)}", numbered=False, size=11)
    body(f"これらを式({n_ab})に代入して，")
    eq(rf"a=\dfrac{{S_{{tE}}}}{{S_t}}=\dfrac{{{tsci(L['Sxy'], 7)}}}{{{tsci(L['Sx'], 7)}}}={tsci(L['a1'], 6)}"
       rf"\ \mathrm{{mV}}/{{}}^{{\circ}}\mathrm{{C}}", numbered=False, size=11)
    eq(rf"b=\bar{{E}}-a\bar{{t}}={tsci(L['yb'], 6)}-{tsci(L['a1'], 6)}\times{tsci(L['xb'], 6)}"
       rf"={tsci(L['b1'], 4)}\ \mathrm{{mV}}", numbered=False, size=11)
    body("を得た．測定値の有効数字を考慮して丸めると，近似直線は")
    n_res_lin = eq(rf"E={tsci(L['a1'], 4)}\,t\ {'+' if L['b1'] >= 0 else '-'}\ {abs(L['b1']):.3f}"
                   rf"\qquad(E\ [\mathrm{{mV}}],\ t\ [{{}}^\circ\mathrm{{C}}])")
    body("となった．")

    para("(ii)　近似曲線（2次式）", sa=2, keep=True)
    body(f"式({n_S2})より，")
    eq(rf"S_{{t^2}}=\overline{{t^4}}-(\overline{{t^2}})^2={tsci(L['x4b'], 7)}-({tsci(L['x2b'], 7)})^2"
       rf"={tsci(L['Sx2'], 7)}", numbered=False, size=11)
    eq(rf"S_{{t^2E}}=\overline{{t^2E}}-\overline{{t^2}}\,\bar{{E}}={tsci(L['x2yb'], 7)}-{tsci(L['x2b'], 7)}"
       rf"\times{tsci(L['yb'], 6)}={tsci(L['Sx2y'], 7)}", numbered=False, size=11)
    eq(rf"S_{{tt^2}}=\overline{{t^3}}-\bar{{t}}\,\overline{{t^2}}={tsci(L['x3b'], 7)}-{tsci(L['xb'], 6)}"
       rf"\times{tsci(L['x2b'], 7)}={tsci(L['Sxx2'], 7)}", numbered=False, size=11)
    body(f"である．S_{{t}}，S_{{tE}} は(i)と同じ値である．式({n_abc})の分母は")
    eq(rf"S_tS_{{t^2}}-S_{{tt^2}}^{{\,2}}={tsci(L['Sx'], 7)}\times{tsci(L['Sx2'], 7)}-({tsci(L['Sxx2'], 7)})^2"
       rf"={tsci(L['Dq'], 7)}", numbered=False, size=10.5)
    numA = L["Sx"] * L["Sx2y"] - L["Sxy"] * L["Sxx2"]
    numB = L["Sxy"] * L["Sx2"] - L["Sxx2"] * L["Sx2y"]
    body("となる．分子はそれぞれ")
    eq(rf"S_tS_{{t^2E}}-S_{{tE}}S_{{tt^2}}={tsci(L['Sx'], 7)}\times{tsci(L['Sx2y'], 7)}"
       rf"-{tsci(L['Sxy'], 7)}\times{tsci(L['Sxx2'], 7)}={tsci(numA, 7)}", numbered=False, size=11)
    eq(rf"S_{{tE}}S_{{t^2}}-S_{{tt^2}}S_{{t^2E}}={tsci(L['Sxy'], 7)}\times{tsci(L['Sx2'], 7)}"
       rf"-{tsci(L['Sxx2'], 7)}\times{tsci(L['Sx2y'], 7)}={tsci(numB, 7)}", numbered=False, size=11)
    body(f"であるから，式({n_abc})より")
    eq(rf"a=\dfrac{{{tsci(numA, 7)}}}{{{tsci(L['Dq'], 7)}}}={tsci(L['a2'], 6)}"
       rf"\ \mathrm{{mV}}/({{}}^{{\circ}}\mathrm{{C}})^2", numbered=False, size=11.5)
    eq(rf"b=\dfrac{{{tsci(numB, 7)}}}{{{tsci(L['Dq'], 7)}}}={tsci(L['b2'], 6)}"
       rf"\ \mathrm{{mV}}/{{}}^{{\circ}}\mathrm{{C}}", numbered=False, size=11.5)
    eq(rf"c=\bar{{E}}-a\overline{{t^2}}-b\bar{{t}}", numbered=False, size=11.5)
    eq(rf"\quad={tsci(L['yb'], 6)}-{tsci(L['a2'], 6)}\times{tsci(L['x2b'], 7)}"
       rf"-{tsci(L['b2'], 6)}\times{tsci(L['xb'], 6)}={tsci(L['c2'], 4)}\ \mathrm{{mV}}", numbered=False, size=11)
    body("を得た．丸めると，近似曲線は")
    sgn = lambda x: "+" if x >= 0 else "-"
    n_res_quad = eq(rf"E={tsci(L['a2'], 3)}\,t^2\ {sgn(L['b2'])}\ {tsci(abs(L['b2']), 4)}\,t\ {sgn(L['c2'])}\ "
                    rf"{abs(L['c2']):.3f}\qquad(E\ [\mathrm{{mV}}],\ t\ [{{}}^\circ\mathrm{{C}}])")
    body("となった．")
    fc = nxt("fig")
    tr = nxt("tab")
    body(f"5つの測定点と，近似直線（式({n_res_lin})）および近似曲線（式({n_res_quad})）を図{fc}に示す．"
         f"また，各点における近似式からの残差（測定値－近似式の値）を表{tr}に示す．")
    figure(os.path.join(IMG, "fig_calibration.png"), "検定に用いた5点と，最小二乗法による近似直線・近似曲線",
           fc, 130)
    rr = []
    for c, r1, r2 in zip(R["comp"], L["res1"], L["res2"]):
        rr.append([c["label"], f"{c['t']:.2f}" if c["t"] else "0", f"{c['E']:.3f}", f"{c['f1']:.3f}",
                   f"{r1:+.3f}", f"{c['f2']:.3f}", f"{r2:+.3f}"])
    rr.append(["残差平方和 Q [mV²]", "", "", "", sci(L['rss1'], 3), "",
               sci(L['rss2'], 3)])
    table("近似式による値と残差（単位：mV）", tr,
          ["定点", "t [℃]", "測定値", "1次式", "残差", "2次式", "残差"], rr, [46, 16, 17, 17, 17, 17, 17],
          size=8.5)
    return {"n_res_lin": n_res_lin, "n_res_quad": n_res_quad, "fig_cal": fc, "tab_res": tr}


refs = draft_results() if not COMPLETE else data_results()

# ================================================================ 6 考察
heading("6　考察")
heading("6.1　残差和や残差の絶対値和を用いてはいけない理由", 2)
body("最小二乗法では残差の二乗和Qを最小にする．残差の和 Σ(y_{i}－ŷ_{i}) を最小（あるいは0）にする"
     "方法が使えないのは，正の残差と負の残差が打ち消し合うためである．1次式の場合，点(x̄, ȳ)を通る"
     "直線であれば傾きによらず残差の和は0になるので，残差の和だけでは直線が一つに定まらない．")
if COMPLETE:
    L = R["lsq"]
    body(f"本実験の5点で確かめると，式({refs['n_res_lin']})の近似直線では残差の和は計算誤差の範囲で0"
         f"（{sci(L['sum_res1'], 2)} mV）であるが，データの傾向をまったく表さない水平な直線 E＝Ē＝{L['yb']:.3f} mV "
         f"も残差の和は同じく0である．一方，残差平方和で比べると，近似直線の{sci(L['rss1'], 3)} mV² に対して"
         f"水平な直線は{L['rss_mean']:.3g} mV² と約{L['rss_mean'] / L['rss1']:.0f}倍も大きく，"
         f"二乗和であれば両者の良し悪しを正しく区別できる．".replace("e-0", "e-"))
body("また，残差の絶対値の和 Σ|y_{i}－ŷ_{i}| は打ち消し合いは起こらないものの，絶対値は残差が0となる点で"
     "微分できないため，Qのように「偏微分して0とおく」方法で係数を求めることができない．係数を求めるには"
     "残差の符号で場合分けしたり，反復計算を行ったりする必要があり，データが増えるほど計算が煩雑になる．"
     "さらに，最小値を与える係数が一つに定まらない場合もある．これに対して二乗和は係数について滑らかな"
     f"下に凸の関数であるため，正規方程式から式({n_ab})，({n_abc})のように係数がただ一つ，簡単な式で求まる．"
     "加えて，測定誤差が正規分布に従う場合，最小二乗法による係数は最も確からしい値（最尤推定値）を"
     "与えるという統計的な根拠もある．ただし，二乗することで大きく外れた点の影響を強く受けるため，"
     "外れ値を含むデータでは絶対値和を最小にする方法（最小絶対値法）が用いられることもある．")

heading("6.2　熱起電力表との比較", 2)
if COMPLETE:
    tc = nxt("tab")
    fd = nxt("fig")
    comp = R["comp"]
    body(f"K熱電対の規格の熱起電力（基準接点0 ℃）[2][3]と，本実験の測定値を比較した結果を表{tc}に示す．"
         f"規格値は，規格の熱起電力表と同じ値を与える基準関数[3]を用いて各定点の温度で計算した．差ΔEは"
         f"その温度におけるゼーベック係数で割って温度のずれΔTにも換算した．なお，0 ℃付近では規格値が"
         f"ほぼ0になるため相対誤差は意味をもたず，ΔEとΔTで評価した．")
    rows = []
    for c in comp:
        rows.append([c["label"], f"{c['t']:.2f}" if c["t"] else "0", f"{c['E']:.3f}", f"{c['Es']:.3f}",
                     f"{c['dE']:+.3f}", f"{c['dT']:+.1f}", "—" if c["rel"] is None else f"{c['rel']:+.1f}",
                     f"±{c['tol']:.1f}"])
    table("測定値と規格の熱起電力の比較", tc,
          ["定点", "t [℃]", "測定値 [mV]", "規格値 [mV]", "ΔE [mV]", "ΔT [℃]", "相対誤差 [%]", "許容差 [℃]"],
          rows, [38, 15, 18, 18, 17, 15, 19, 18], size=8.5)
    dEs = [c["dE"] for c in comp]
    dTs = [c["dT"] for c in comp]
    melts = comp[2:]
    allneg = all(d < 0 for d in dEs)
    allpos = all(d > 0 for d in dEs)
    within = all(abs(c["dT"]) <= c["tol"] for c in comp)
    s = (f"表{tc}より，測定値と規格値の差は {min(dEs):+.3f}〜{max(dEs):+.3f} mV，温度に換算すると "
         f"{min(dTs):+.1f}〜{max(dTs):+.1f} ℃ であった．")
    nz = dEs[1:]   # 0 ℃以外（室温・融点）
    if allneg or allpos:
        s += (f"すべての定点で測定値は規格値より{'小さ' if allneg else '大き'}く，温度を"
              f"{'低' if allneg else '高'}めに示す系統的な傾向が見られた．")
    elif all(d < 0 for d in nz) or all(d > 0 for d in nz):
        neg = all(d < 0 for d in nz)
        s += (f"0 ℃の点を除くすべての定点で測定値は規格値より{'小さ' if neg else '大き'}く，温度を"
              f"{'低' if neg else '高'}めに示す系統的な傾向が見られた．また，そのずれは温度が高い定点ほど"
              f"大きくなる傾向にあった．" if all(abs(nz[i]) <= abs(nz[i + 1]) for i in range(len(nz) - 1)) else
              f"0 ℃の点を除くすべての定点で測定値は規格値より{'小さ' if neg else '大き'}く，温度を"
              f"{'低' if neg else '高'}めに示す系統的な傾向が見られた．")
    else:
        s += "測定値が規格値より大きい点と小さい点があり，一方向の系統的なずれは見られなかった．"
    s += (f"金属の融点における相対誤差は{min(abs(c['rel']) for c in melts):.1f}〜"
          f"{max(abs(c['rel']) for c in melts):.1f} %であった．")
    s += ("いずれの定点でも，ずれはK熱電対のクラス2の許容差（±2.5 ℃または±0.0075|t|の大きい方）[2]"
          "の範囲内であり，本実験の熱電対は規格の熱起電力表に従って温度計として使用できる精度をもつと"
          "いえる．" if within else
          "一部の定点ではずれがK熱電対のクラス2の許容差（±2.5 ℃または±0.0075|t|の大きい方）[2]を"
          "超えており，この熱電対を規格の熱起電力表のまま用いると誤差が大きくなるため，本実験で求めた"
          "検定式を用いて補正する意義があるといえる．")
    body(s)
    z = D.ZERO_EMF
    corr = [c["dE0"] for c in comp[1:]]
    body(f"ずれの原因を検討する．まず，0 ℃で測定した起電力 E₀＝{z:.3f} mV は理想値0 mVからずれており，"
         f"これは基準接点の温度がわずかに0 ℃からずれていたことや，端子部分で生じる微小な熱起電力，"
         f"ボルトメータのゼロ点のずれなど，全測定に共通に加わるオフセットと考えられる．各定点のΔEから"
         f"E₀を差し引くと，室温および融点におけるずれは {min(corr):+.3f}〜{max(corr):+.3f} mV となり，"
         + ("オフセットでずれの一部を説明できる．" if max(abs(x) for x in corr) < max(abs(x) for x in dEs[1:])
            else "オフセットだけではずれを説明できない．"))
    body("次に，融点の測定では，熱電対の先端が保護管とタンマン管の壁を隔てて試料に接しているため，素線や"
         "保護管を伝って熱が外へ逃げ，測温接点の温度が試料の温度よりわずかに低くなることが考えられる．"
         "また，試料に不純物が含まれると凝固点降下によって凝固点が下がり，プラトーの区間の平均には"
         "温度がゆるやかに下がり始めた部分も含まれる．これらはいずれも起電力を規格値より小さくする方向に"
         "働く．室温については，室内の温度計の位置と測温接点の位置で温度が異なる可能性があり，"
         "定点の温度そのものに±1 ℃程度の不確かさがあると考えられる．")
    body(f"図{fd}に，近似直線・近似曲線と測定点のそれぞれについて，規格の熱起電力からのずれを温度に換算して"
         f"示す．")
    figure(os.path.join(IMG, "fig_deviation.png"), "規格の熱起電力からのずれ（温度換算）", fd, 125)
else:
    todo("【測定データ受領後に作成】JIS C 1602の熱起電力表（基準接点0 ℃）の値と測定値を比較する表（ΔE，温度換算"
         "ΔT，相対誤差，クラス2許容差との比較），ずれの傾向と原因（0 ℃のオフセット，熱の逃げ，凝固点降下，"
         "プラトーの決め方）の定量的な検討．")

heading("6.3　直線および2次曲線による検定の結果から分かること", 2)
if COMPLETE:
    L = R["lsq"]
    comp = R["comp"]
    dev = R["dev"]
    std_a = R["std_quad"][0]
    errs_lin = [c["t_lin"] - c["t"] for c in comp]
    errs_quad = [c["t_quad"] - c["t"] for c in comp]
    body(f"表{refs['tab_res']}より，残差平方和Qは1次式で{sci(L['rss1'], 3)} mV²，2次式で{sci(L['rss2'], 3)} mV²であり，"
         f"2次式の方が小さかった．ただし，2次式は1次式に係数を一つ加えた式で，a＝0とすれば1次式に戻るので，"
         f"同じデータに当てはめれば2次式のQが1次式より大きくなることはない．とくに本実験は5点に対して"
         f"係数が3個であり，自由度が小さいため，Qが小さいことだけでは2次式の方が真の関係に近いとは"
         f"いえない．".replace("e-0", "e-"))
    same = (L["a2"] > 0) == (std_a > 0)
    body(f"そこで，2次の係数の意味を規格値と比べて検討する．規格の熱起電力を同じ5つの温度で2次式に"
         f"近似すると，2次の係数は{sci(std_a, 3)} mV/℃² となる．本実験の2次の係数 {sci(L['a2'], 3)} mV/℃² は"
         + (f"これと符号が同じで，大きさも{abs(L['a2'] / std_a):.1f}倍と同程度である．したがって，2次式で表された"
            f"わずかな曲がりは測定のばらつきではなく，K熱電対のゼーベック係数が温度とともに増加するという"
            f"実際の性質を反映していると考えられる．" if same and 0.3 < abs(L['a2'] / std_a) < 3 else
            f"これと{'符号は同じだが大きさが大きく異なる' if same else '符号が異なる'}．したがって，本実験の2次の"
            f"係数には測定のばらつきの影響が大きく含まれていると考えられる．"))
    lm = max(abs(e) for e in errs_lin)
    qm = max(abs(e) for e in errs_quad)
    jl = max(range(len(comp)), key=lambda i: abs(errs_lin[i]))
    txt = (f"また，求めた検定式を温度計として用いる場合を考え，各定点の測定起電力を検定式で温度に戻したときの"
           f"誤差を求めると，最大で1次式では{lm:.1f} ℃，2次式では{qm:.1f} ℃であった．")
    if lm < 2.5 and qm < 2.5:
        txt += ("いずれもK熱電対クラス2の許容差（±2.5 ℃）[2]より小さく，0〜420 ℃の範囲では1次式でも"
                "実用上十分な精度で温度を求められることが分かる．")
    elif qm < 2.5:
        txt += (f"1次式では{comp[jl]['label']}で{errs_lin[jl]:+.1f} ℃となり，K熱電対クラス2の許容差（±2.5 ℃）[2]"
                f"を超えたのに対し，2次式ではすべての定点で許容差以内に収まった．直線は温度範囲の両端付近で"
                f"実際の曲がった関係から離れやすく，とくに低温側でずれが目立つためである．")
    else:
        txt += "どちらの式でも許容差（±2.5 ℃）[2]を超える定点があり，定点数を増やした検定が必要である．"
    body(txt)
    concl = (f"一方，図{fd}のように0〜450 ℃の範囲で規格の熱起電力と比べると，最大のずれ（温度換算）は1次式で"
             f"{dev['max1_C']:.1f} ℃，2次式で{dev['max2_C']:.1f} ℃であった．検定式は測定点に合わせて決めているため，"
             f"規格値とのずれには6.2で述べた測定の系統的なずれも含まれている．また，2次式は定点の間を滑らかに"
             f"つなぐ反面，測定した範囲（0〜420 ℃）の外側では2次の項が急に効いてくるため，範囲外へ外挿して"
             f"使うのは危険である．以上より，")
    if lm < 2.5:
        concl += ("本実験の温度範囲では1次式でも十分な精度が得られるが，ゼーベック係数の温度変化まで表すには"
                  "2次式が適している．")
    else:
        concl += ("K熱電対の起電力と温度の関係には無視できない曲がりがあり，0〜420 ℃の範囲を一つの式で"
                  "表すには2次式が適している．1次式は平均的なゼーベック係数を知るには便利であるが，温度範囲の"
                  "端，とくに低温側で誤差が大きくなる．")
    concl += "また，より正確な検定を行うには，定点の数を増やすことが望ましい．"
    body(concl)
else:
    todo("【測定データ受領後に作成】1次式と2次式の残差平方和の比較（2次式は1次式を含むためQが小さくなるのは当然で"
         "あること），2次の係数と規格値の曲がりとの比較，検定式で起電力を温度に戻したときの誤差，測定範囲外への"
         "外挿の危険性．")

heading("6.4　ゼーベック係数の推定", 2)
if COMPLETE:
    sb_ = R["seebeck"]
    L = R["lsq"]
    q = sb_["quad"]
    s = sb_["std"]
    body(f"熱電対としてのゼーベック係数は S_{{AB}}＝dE/dt で与えられる（式({n_seebeck})）．近似直線の傾きから，"
         f"0〜420 ℃の平均的なゼーベック係数は")
    eq(rf"S_{{AB}}=a={tsci(L['a1'], 4)}\ \mathrm{{mV}}/{{}}^{{\circ}}\mathrm{{C}}={L['a1'] * 1000:.1f}\ \mu\mathrm{{V/K}}",
       numbered=False, size=11.5)
    body(f"と推定される（温度差1 ℃と1 Kは等しいので，単位はµV/Kとしても同じ値である）．規格値から求めた"
         f"同じ温度範囲の平均の傾き {(sb_['std_avg']) * 1000:.1f} µV/K[3] と比べると，差は"
         f"{(L['a1'] - sb_['std_avg']) / sb_['std_avg'] * 100:+.1f} %である．")
    body(f"また，近似曲線（式({refs['n_res_quad']})）を微分すると S_{{AB}}＝2at＋b となり，温度によって変化する"
         f"ゼーベック係数が得られる．0 ℃，200 ℃，400 ℃ではそれぞれ {q[0] * 1000:.1f}，{q[200] * 1000:.1f}，"
         f"{q[400] * 1000:.1f} µV/K となり，規格値の {s[0] * 1000:.1f}，{s[200] * 1000:.1f}，{s[400] * 1000:.1f} µV/K "
         f"[3] と同様に，温度が高いほど大きくなる傾向が得られた．"
         + ("" if q[400] > q[0] else "ただし，本実験の結果は温度とともに減少しており，規格値とは傾向が異なる．"))
    body("なお，ここで求めたのはクロメルとアルメルのゼーベック係数の差（相対ゼーベック係数）であり，"
         "正の値は高温側の接点から見てクロメル側が＋極になることに対応する．")
else:
    todo("【測定データ受領後に作成】近似直線の傾きから平均のゼーベック係数（µV/K），近似曲線の微分 2at＋b から"
         "0・200・400 ℃での値を求め，規格値と比較する．")

if COMPLETE and all("heat" in R["cool"][m] for m in A.ORDER):
    heading("6.5　加熱時と冷却時のプラトーの比較", 2)
    th = nxt("tab")
    rows = []
    diffs = []
    for m in A.ORDER:
        c = R["cool"][m]
        h = c["heat"]
        d = h["mean"] - c["mean"]
        diffs.append(d)
        rows.append([f"{A.NAME_JP[m]}（{m}）", f"{h['t0']:.0f}〜{h['t1']:.0f}", f"{h['mean']:.3f}",
                     f"{h['sd']:.3f}", f"{c['mean']:.3f}", f"{c['sd']:.3f}", f"{d:+.3f}",
                     f"{d / A.seebeck_std(A.MELT[m]):+.1f}"])
    body(f"本実験では加熱開始から通して記録したため，加熱中にも熱起電力の上昇が鈍る区間が観測された．"
         f"隣り合う測定値の差が{A.HEAT_THR:.2f} mV以下の区間を加熱時のプラトーとして，その平均を冷却時の"
         f"プラトー（表5）と比べた結果を表{th}に示す．加熱時の区間の時刻は加熱開始からの時刻である．")
    table("加熱時と冷却時のプラトーの比較", th,
          ["金属", "加熱時の区間 [s]", "加熱時 平均 [mV]", "標準偏差 [mV]", "冷却時 平均 [mV]",
           "標準偏差 [mV]", "差 [mV]", "差 [℃]"], rows, [24, 22, 19, 17, 19, 17, 15, 14], size=8)
    hi = all(d > 0 for d in diffs)
    body(f"表{th}より，加熱時のプラトーの平均は冷却時に比べて{min(diffs):+.3f}〜{max(diffs):+.3f} mV"
         + ("となり，すべての金属で加熱時の方が高かった．" if hi else "の差があった．")
         + "加熱中は電気炉から試料へ熱が流れ込み続けるため，炉壁に近い外側から融解が進む一方で，試料の"
         "内部や保護管の中の熱電対との間に温度差が残り，熱電対は融点より高めの温度を示しやすい．"
         + ("また，加熱時の区間は標準偏差が冷却時より大きく，起電力が一定にならずにゆるやかに上昇し続けて"
            "いる．" if all(R["cool"][m]["heat"]["sd"] > R["cool"][m]["sd"] for m in A.ORDER) else "")
         + "これに対して冷却時は通電を止めて温度がゆっくり下がるため試料内の温度差が小さく，凝固の間は"
         "潜熱によって温度が一定に保たれる．以上の比較から，融点に対応する起電力を冷却時のプラトーから"
         "決めるという本実験の方法が妥当であることを，自分の測定データで確かめることができた．")

# ================================================================ 参考文献
heading("参考文献")
refs_list = [
    "[1] 東京理科大学，“物理学実験 実験指導書 15．熱電対の検定実験，”2025．",
    "[2] 日本産業標準調査会，“JIS C 1602:2015 熱電対，”日本規格協会，2015．",
    "[3] G. W. Burns, M. G. Scroger, G. F. Strouse, M. C. Croarkin, W. F. Guthrie, “Temperature-"
    "Electromotive Force Reference Functions and Tables for the Letter-Designated Thermocouple Types "
    "Based on the ITS-90,” NIST Monograph 175, National Institute of Standards and Technology, 1993.",
]
for r in refs_list:
    p = para(r, size=10, sa=3, line=1.3)
    p.paragraph_format.left_indent = Mm(7)
    p.paragraph_format.first_line_indent = Mm(-7)

# ---------------------------------------------------------------- フッタ（ページ番号）
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run()
for tag, attr, val in [("w:fldChar", "w:fldCharType", "begin"), ("w:instrText", None, "PAGE"),
                       ("w:fldChar", "w:fldCharType", "end")]:
    el = OxmlElement(tag)
    if attr:
        el.set(qn(attr), val)
    else:
        el.set(qn("xml:space"), "preserve")
        el.text = val
    run._r.append(el)
font(run, size=10)
sec.different_first_page_header_footer = True

doc.save(out_path)
print("saved:", out_path, "| complete data:", COMPLETE)
