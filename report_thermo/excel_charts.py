# -*- coding: utf-8 -*-
"""熱電対の検定実験 — Excel ブック（データ・最小二乗法の数式・グラフ）の作成と、
Excel グラフを報告書用の画像にする処理。

- build_workbook(): 提出・確認用の .xlsx（Excel ネイティブのグラフ、計算は数式）
- render_figures(): 同じ定義のグラフを LibreOffice で PDF 化し、PNG に切り出して img/ に保存
"""
import io
import os
import re
import shutil
import zipfile
import subprocess
import tempfile

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from lxml import etree
from PIL import Image, ImageChops

import analysis as A

FONT_JP = "游ゴシック"
BLUE, ORANGE, GRAY = "4472C4", "ED7D31", "A5A5A5"   # Excel 既定の系列色（Office テーマ）
AXIS_GRAY, TEXT_GRAY = "D9D9D9", "595959"
THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ---------------------------------------------------------------- グラフ部品
def _marker(sym, size, fill="000000", edge="000000"):
    return Marker(symbol=sym, size=size,
                  spPr=GraphicalProperties(solidFill=fill, ln=LineProperties(solidFill=edge, w=9525)))


def _line(series, color="000000", width=9525, dash=None, none=False):
    gp = series.graphicalProperties
    if none:
        gp.line.noFill = True
        return
    gp.line.solidFill = color
    gp.line.width = width
    if dash:
        gp.line.dashStyle = dash


def _base_chart(xtitle, ytitle, w=16, h=9.5):
    ch = ScatterChart()
    ch.style = 2
    ch.title = None
    ch.x_axis.title = xtitle
    ch.y_axis.title = ytitle
    for ax in (ch.x_axis, ch.y_axis):
        ax.delete = False
        ax.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill=AXIS_GRAY, w=9525)))
        ax.majorTickMark = "none"
        ax.minorTickMark = "none"
        ax.spPr = GraphicalProperties(ln=LineProperties(solidFill=AXIS_GRAY, w=9525))
        ax.number_format = "General"
    ch.legend.position = "b"
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.width, ch.height = w, h
    return ch


def _nice(lo, hi, step):
    import math
    return math.floor(lo / step) * step, math.ceil(hi / step) * step


# ---------------------------------------------------------------- シート作成
def _header(ws, row, headers, col=1):
    for j, h in enumerate(headers):
        c = ws.cell(row=row, column=col + j, value=h)
        c.font = Font(name=FONT_JP, bold=True)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = BOX


def _cell(ws, r, c, v, fmt=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.border = BOX
    cell.font = Font(name=FONT_JP)
    if fmt:
        cell.number_format = fmt
    return cell


def sheet_raw(wb, D):
    ws = wb.create_sheet("測定データ")
    ws["A1"] = "実験3：加熱開始を0 sとした通しの記録（15 s間隔）。起電力が最大の点以降を冷却時のデータとした．"
    ws["A1"].font = Font(name=FONT_JP)
    col = 1
    for m in A.ORDER:
        E = D.RAW[m]
        ws.cell(row=3, column=col, value=f"{A.NAME_JP[m]}（{m}）").font = Font(name=FONT_JP, bold=True)
        _header(ws, 4, ["時間 [s]", "熱起電力 [mV]"], col)
        for i, e in enumerate(E):
            _cell(ws, 5 + i, col, i * D.DT, "0")
            _cell(ws, 5 + i, col + 1, e, "0.000")
        col += 3
    ws.cell(row=3, column=col, value="実験1・2").font = Font(name=FONT_JP, bold=True)
    _header(ws, 4, ["項目", "温度 [℃]", "熱起電力 [mV]"], col)
    _cell(ws, 5, col, "室温"); _cell(ws, 5, col + 1, D.ROOM_TEMP, "0.0"); _cell(ws, 5, col + 2, D.ROOM_EMF, "0.000")
    _cell(ws, 6, col, "0 ℃（氷水）"); _cell(ws, 6, col + 1, 0, "0"); _cell(ws, 6, col + 2, D.ZERO_EMF, "0.000")
    for c in range(1, col + 3):
        ws.column_dimensions[get_column_letter(c)].width = 13
    return ws


def sheet_cooling(wb, R, D, m, anchor_chart=True):
    """冷却データ＋プラトーの平均・標準偏差（数式）＋冷却曲線グラフ。"""
    c = R["cool"][m]
    ws = wb.create_sheet(f"冷却_{m}")
    ws["A1"] = (f"{A.NAME_JP[m]}（{m}，融点 {A.MELT[m]:.2f} ℃）の冷却データ"
                f"（起電力が最大となった時点を0 sとした）")
    ws["A1"].font = Font(name=FONT_JP, bold=True)
    _header(ws, 3, ["時間 [s]", "熱起電力 [mV]", "融点とみなした区間 [mV]"])
    n = len(c["E"])
    for i, (t, e) in enumerate(zip(c["t"], c["E"])):
        r = 4 + i
        _cell(ws, r, 1, t, "0")
        _cell(ws, r, 2, e, "0.000")
        _cell(ws, r, 3, e if c["i0"] <= i <= c["i1"] else None, "0.000")
    r0, r1 = 4 + c["i0"], 4 + c["i1"]
    # プラトーの統計（数式）
    _header(ws, 3, ["プラトーの統計", "値"], 5)
    rows = [("区間 [s]", f"{c['t0']:.0f}〜{c['t1']:.0f}", None),
            ("点数", f"=COUNT(C{r0}:C{r1})", "0"),
            ("平均 [mV]", f"=AVERAGE(C{r0}:C{r1})", "0.000"),
            ("標準偏差 [mV]", f"=STDEV(C{r0}:C{r1})", "0.000"),
            ("（計算値）平均 [mV]", round(c["mean"], 6), "0.000")]
    for k, (lab, v, f) in enumerate(rows):
        _cell(ws, 4 + k, 5, lab)
        _cell(ws, 4 + k, 6, v, f)
    ws["E10"] = f"判定基準：隣り合う測定値の差が {A.PLATEAU_THR:.2f} mV 以下の状態が最も長く続いた区間"
    ws["E10"].font = Font(name=FONT_JP, size=9)
    # 平均線（グラフ用）
    _header(ws, 12, ["平均線 x [s]", "平均線 y [mV]"], 5)
    _cell(ws, 13, 5, 0); _cell(ws, 13, 6, round(c["mean"], 6), "0.000")
    _cell(ws, 14, 5, c["t"][-1]); _cell(ws, 14, 6, round(c["mean"], 6), "0.000")
    for col, w in zip("ABCDEF", [10, 14, 20, 3, 20, 14]):
        ws.column_dimensions[col].width = w

    ch = _base_chart("時間 [s]", "熱起電力 [mV]")
    x = Reference(ws, min_col=1, min_row=4, max_row=3 + n)
    s1 = Series(Reference(ws, min_col=2, min_row=4, max_row=3 + n), x, title="測定値")
    s1.marker = _marker("circle", 5, BLUE, BLUE)
    _line(s1, BLUE, width=28575)
    s2 = Series(Reference(ws, min_col=3, min_row=4, max_row=3 + n), x, title="融点とみなした区間")
    s2.marker = _marker("circle", 5, ORANGE, ORANGE)
    _line(s2, none=True)
    s3 = Series(Reference(ws, min_col=6, min_row=13, max_row=14), Reference(ws, min_col=5, min_row=13, max_row=14),
                title=f"平均値 {c['mean']:.3f} mV")
    s3.marker = Marker(symbol="none")
    _line(s3, GRAY, width=19050, dash="dash")
    for s in (s1, s2, s3):
        ch.series.append(s)
    lo, hi = _nice(min(c["E"]) - 0.1, max(c["E"]) + 0.1, 0.5)
    ch.y_axis.scaling.min, ch.y_axis.scaling.max = lo, hi
    ch.y_axis.majorUnit = 0.5 if hi - lo <= 3 else 1.0
    xmax = _nice(0, c["t"][-1], 100)[1]
    ch.x_axis.scaling.min, ch.x_axis.scaling.max = 0, xmax
    ch.x_axis.majorUnit = 100 if xmax <= 800 else 200
    if anchor_chart:
        ws.add_chart(ch, "H3")
    return ws, ch


def sheet_lsq(wb, R, anchor_chart=True):
    """最小二乗法の計算（すべて数式）と，5点・近似直線・近似曲線のグラフ。"""
    L = R["lsq"]
    ws = wb.create_sheet("最小二乗法")
    ws["A1"] = "熱電対の検定：最小二乗法（x：温度 t [℃]，y：熱起電力 E [mV]，N＝5）"
    ws["A1"].font = Font(name=FONT_JP, bold=True)
    _header(ws, 3, ["定点", "t", "E", "t²", "tE", "t³", "t⁴", "t²E"])
    for i, (lab, t, e) in enumerate(R["points"]):
        r = 4 + i
        _cell(ws, r, 1, lab)
        _cell(ws, r, 2, t, "0.00")
        _cell(ws, r, 3, e, "0.000")
        _cell(ws, r, 4, f"=B{r}^2", "0.000E+00")
        _cell(ws, r, 5, f"=B{r}*C{r}", "0.000E+00")
        _cell(ws, r, 6, f"=B{r}^3", "0.000E+00")
        _cell(ws, r, 7, f"=B{r}^4", "0.000E+00")
        _cell(ws, r, 8, f"=B{r}^2*C{r}", "0.000E+00")
    _cell(ws, 9, 1, "平均")
    for col in "BCDEFGH":
        _cell(ws, 9, ord(col) - 64, f"=AVERAGE({col}4:{col}8)", "0.000000E+00")
    calc = [
        ("S_t（tの分散）", "=D9-B9^2"),
        ("S_tE（tとEの共分散）", "=E9-B9*C9"),
        ("S_t²（t²の分散）", "=G9-D9^2"),
        ("S_t²E（t²とEの共分散）", "=H9-D9*C9"),
        ("S_tt²（tとt²の共分散）", "=F9-B9*D9"),
        ("1次式 a [mV/℃]", "=B13/B12"),
        ("1次式 b [mV]", "=C9-B17*B9"),
        ("分母 S_t·S_t² − S_tt²²", "=B12*B14-B16^2"),
        ("2次式 a [mV/℃²]", "=(B12*B15-B13*B16)/B19"),
        ("2次式 b [mV/℃]", "=(B13*B14-B16*B15)/B19"),
        ("2次式 c [mV]", "=C9-B20*D9-B21*B9"),
        ("残差平方和 Q（1次式）[mV²]", "=SUMPRODUCT((C4:C8-(B17*B4:B8+B18))^2)"),
        ("残差平方和 Q（2次式）[mV²]", "=SUMPRODUCT((C4:C8-(B20*B4:B8^2+B21*B4:B8+B22))^2)"),
    ]
    _header(ws, 11, ["量", "値（数式）", "検算（Python）"])
    py = [L["Sx"], L["Sxy"], L["Sx2"], L["Sx2y"], L["Sxx2"], L["a1"], L["b1"], L["Dq"],
          L["a2"], L["b2"], L["c2"], L["rss1"], L["rss2"]]
    for k, ((lab, f), pv) in enumerate(zip(calc, py)):
        _cell(ws, 12 + k, 1, lab)
        _cell(ws, 12 + k, 2, f, "0.000000E+00")
        _cell(ws, 12 + k, 3, pv, "0.000000E+00")
    ws["A26"] = ("※ 係数は Excel の近似曲線機能ではなく，テキストの式(7)〜(16)をそのまま数式にして求めた．"
                 "C列は報告書作成時の計算値（同じ値になることの確認用）．")
    ws["A26"].font = Font(name=FONT_JP, size=9)
    # 近似式の値（グラフ用，数式）
    _header(ws, 3, ["t [℃]", "1次式 E [mV]", "2次式 E [mV]"], 11)
    ts = list(range(0, 451, 10))
    for i, t in enumerate(ts):
        r = 4 + i
        _cell(ws, r, 11, t, "0")
        # LibreOffice での描画用に計算値も入れるため，数式ではなく値を入れ，式は注記で示す
        _cell(ws, r, 12, L["a1"] * t + L["b1"], "0.000")
        _cell(ws, r, 13, L["a2"] * t * t + L["b2"] * t + L["c2"], "0.000")
    ws.cell(row=2, column=11, value="近似式の値（L列＝$B$17·t＋$B$18，M列＝$B$20·t²＋$B$21·t＋$B$22）").font = \
        Font(name=FONT_JP, size=9)
    for col, w in zip("ABCDEFGH", [30, 16, 16, 13, 13, 13, 13, 13]):
        ws.column_dimensions[col].width = w

    ch = _base_chart("温度 t [℃]", "熱起電力 E [mV]", w=16, h=11)
    nx = len(ts)
    xs = Reference(ws, min_col=11, min_row=4, max_row=3 + nx)
    s_l = Series(Reference(ws, min_col=12, min_row=4, max_row=3 + nx), xs, title="近似直線（1次式）")
    s_l.marker = Marker(symbol="none"); _line(s_l, ORANGE, width=28575)
    s_q = Series(Reference(ws, min_col=13, min_row=4, max_row=3 + nx), xs, title="近似曲線（2次式）")
    s_q.marker = Marker(symbol="none"); _line(s_q, GRAY, width=28575, dash="dash")
    s_p = Series(Reference(ws, min_col=3, min_row=4, max_row=8), Reference(ws, min_col=2, min_row=4, max_row=8),
                 title="測定点")
    s_p.marker = _marker("circle", 7, BLUE, BLUE); _line(s_p, none=True)
    for s in (s_l, s_q, s_p):
        ch.series.append(s)
    ch.x_axis.scaling.min, ch.x_axis.scaling.max, ch.x_axis.majorUnit = 0, 450, 50
    ch.y_axis.scaling.min, ch.y_axis.scaling.max, ch.y_axis.majorUnit = 0, 20, 2
    if anchor_chart:
        ws.add_chart(ch, "O3")
    return ws, ch


def sheet_dev(wb, R, anchor_chart=True):
    """規格の熱起電力（JIS C 1602 / NIST）からのずれ（温度換算）。"""
    L = R["lsq"]
    ws = wb.create_sheet("規格との比較")
    ws["A1"] = "規格の熱起電力（基準接点0 ℃）との比較．規格値は JIS C 1602 の熱起電力表と同じ値を与える基準関数で計算．"
    ws["A1"].font = Font(name=FONT_JP, bold=True)
    _header(ws, 3, ["定点", "t [℃]", "測定値 [mV]", "規格値 [mV]", "ΔE [mV]", "ΔT [℃]", "許容差 [℃]"])
    for i, c in enumerate(R["comp"]):
        r = 4 + i
        _cell(ws, r, 1, c["label"]); _cell(ws, r, 2, c["t"], "0.00"); _cell(ws, r, 3, c["E"], "0.000")
        _cell(ws, r, 4, c["Es"], "0.000"); _cell(ws, r, 5, f"=C{r}-D{r}", "+0.000;-0.000")
        _cell(ws, r, 6, round(c["dT"], 4), "+0.0;-0.0"); _cell(ws, r, 7, c["tol"], "0.0")
    _header(ws, 11, ["t [℃]", "1次式のずれ [℃]", "2次式のずれ [℃]"])
    ts = list(range(0, 451, 5))
    for i, t in enumerate(ts):
        r = 12 + i
        s = A.seebeck_std(t)
        std = A.emf_std(t)
        _cell(ws, r, 1, t, "0")
        _cell(ws, r, 2, ((L["a1"] * t + L["b1"]) - std) / s, "0.00")
        _cell(ws, r, 3, ((L["a2"] * t * t + L["b2"] * t + L["c2"]) - std) / s, "0.00")
    for col, w in zip("ABCDEFG", [18, 12, 14, 14, 12, 12, 12]):
        ws.column_dimensions[col].width = w

    ch = _base_chart("温度 t [℃]", "規格からのずれ（温度換算）[℃]", w=16, h=9.5)
    n = len(ts)
    xs = Reference(ws, min_col=1, min_row=12, max_row=11 + n)
    s1 = Series(Reference(ws, min_col=2, min_row=12, max_row=11 + n), xs, title="近似直線（1次式）")
    s1.marker = Marker(symbol="none"); _line(s1, ORANGE, width=28575)
    s2 = Series(Reference(ws, min_col=3, min_row=12, max_row=11 + n), xs, title="近似曲線（2次式）")
    s2.marker = Marker(symbol="none"); _line(s2, GRAY, width=28575, dash="dash")
    s3 = Series(Reference(ws, min_col=6, min_row=4, max_row=8), Reference(ws, min_col=2, min_row=4, max_row=8),
                title="測定点")
    s3.marker = _marker("circle", 7, BLUE, BLUE); _line(s3, none=True)
    for s in (s1, s2, s3):
        ch.series.append(s)
    vals = [x for c in R["comp"] for x in [c["dT"]]]
    import numpy as np
    d1 = [((L["a1"] * t + L["b1"]) - A.emf_std(t)) / A.seebeck_std(t) for t in ts]
    d2 = [((L["a2"] * t * t + L["b2"] * t + L["c2"]) - A.emf_std(t)) / A.seebeck_std(t) for t in ts]
    lo, hi = _nice(min(vals + d1 + d2) - 0.5, max(vals + d1 + d2 + [0]) + 0.5, 1)
    ch.y_axis.scaling.min, ch.y_axis.scaling.max, ch.y_axis.majorUnit = lo, hi, 1
    ch.x_axis.scaling.min, ch.x_axis.scaling.max, ch.x_axis.majorUnit = 0, 450, 50
    ch.x_axis.crosses = "min"
    if anchor_chart:
        ws.add_chart(ch, "I3")
    return ws, ch


# ---------------------------------------------------------------- 提出・確認用ブック
def build_workbook(R, D, path):
    wb = Workbook()
    wb.remove(wb.active)
    sheet_raw(wb, D)
    for m in A.ORDER:
        sheet_cooling(wb, R, D, m)
    sheet_lsq(wb, R)
    sheet_dev(wb, R)
    wb.calculation.fullCalcOnLoad = True
    bio = io.BytesIO()
    wb.save(bio)
    with open(path, "wb") as f:
        f.write(_patch_xlsx(bio.getvalue()))
    return path


# ---------------------------------------------------------------- 報告書用の画像
FIG_NAMES = [f"fig_cool_{m}" for m in A.ORDER] + ["fig_calibration", "fig_deviation"]


def _trim(im, pad=6):
    bg = Image.new(im.mode, im.size, (255, 255, 255))
    box = ImageChops.difference(im, bg).getbbox()
    if not box:
        return im
    l, t, r, b = box
    return im.crop((max(l - pad, 0), max(t - pad, 0), min(r + pad, im.width), min(b + pad, im.height)))


def render_figures(R, D, imgdir, dpi=220):
    """グラフだけを並べたブックを作り，LibreOffice で PDF にして 1 ページ 1 図で切り出す。"""
    wb = Workbook()
    wb.remove(wb.active)
    charts = []
    for m in A.ORDER:
        charts.append(sheet_cooling(wb, R, D, m, anchor_chart=False))
    charts.append(sheet_lsq(wb, R, anchor_chart=False))
    charts.append(sheet_dev(wb, R, anchor_chart=False))
    for (ws, _), name in zip(charts, FIG_NAMES):
        ws.sheet_state = "hidden"
    for (ws, ch), name in zip(charts, FIG_NAMES):
        g = wb.create_sheet(name)
        g.add_chart(ch, "A1")
    wb.active = len(wb.sheetnames) - len(FIG_NAMES)
    tmp = tempfile.mkdtemp()
    x = os.path.join(tmp, "figs.xlsx")
    wb.save(x)
    prof = os.path.join(tmp, "lo")
    subprocess.run(["soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--convert-to", "pdf",
                    x, "--outdir", tmp], check=True, capture_output=True, timeout=300)
    pdf = os.path.join(tmp, "figs.pdf")
    subprocess.run(["pdftoppm", "-r", str(dpi), "-png", pdf, os.path.join(tmp, "p")], check=True)
    pages = sorted(f for f in os.listdir(tmp) if f.startswith("p-") and f.endswith(".png"))
    if len(pages) != len(FIG_NAMES):
        raise RuntimeError(f"グラフのページ数が一致しません: {len(pages)}")
    os.makedirs(imgdir, exist_ok=True)
    for pg, name in zip(pages, FIG_NAMES):
        im = Image.open(os.path.join(tmp, pg)).convert("RGB")
        _trim(im).save(os.path.join(imgdir, name + ".png"))
    shutil.rmtree(tmp, ignore_errors=True)


# ---------------------------------------------------------------- Word に埋め込む本物のグラフ
import io
import re
import zipfile

from lxml import etree

C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _ref_values(ws, formula):
    """'シート名'!$B$4:$B$42 の値のリスト（空欄は None）。"""
    m = re.match(r"^'?(.*?)'?!\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+)$", formula)
    if not m:
        raise ValueError(formula)
    col, r0, r1 = m.group(2), int(m.group(3)), int(m.group(5))
    return [ws[f"{col}{r}"].value for r in range(r0, r1 + 1)]


A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _q(t):
    return f"{{{C_NS}}}{t}"


def _fix_layout(root):
    """Excel が保存するときと同じ指定を加える：軸の題名と凡例をグラフ領域に重ねない，
    グラフタイトルなし，角丸なし，文字は Excel 既定と同じ游ゴシック・濃い灰色・9 pt。"""
    if root.find(_q("roundedCorners")) is None:
        rc = etree.Element(_q("roundedCorners"))
        rc.set("val", "0")
        root.insert(0, rc)
    for title in root.iter(_q("title")):
        if title.find(_q("overlay")) is None:
            ov = etree.Element(_q("overlay"))
            ov.set("val", "0")
            title.find(_q("tx")).addnext(ov)
    for legend in root.iter(_q("legend")):
        if legend.find(_q("overlay")) is None:
            ov = etree.Element(_q("overlay"))
            ov.set("val", "0")
            legend.find(_q("legendPos")).addnext(ov)
    chart = root.find(_q("chart"))
    if chart.find(_q("autoTitleDeleted")) is None:
        atd = etree.Element(_q("autoTitleDeleted"))
        atd.set("val", "1")
        chart.insert(0, atd)
    if root.find(_q("txPr")) is None:
        tx = etree.fromstring(
            f'<c:txPr xmlns:c="{C_NS}" xmlns:a="{A_NS}"><a:bodyPr/><a:lstStyle/><a:p><a:pPr>'
            f'<a:defRPr sz="900"><a:solidFill><a:srgbClr val="{TEXT_GRAY}"/></a:solidFill>'
            f'<a:latin typeface="游ゴシック"/><a:ea typeface="游ゴシック"/>'
            f'</a:defRPr></a:pPr><a:endParaRPr lang="ja-JP"/></a:p></c:txPr>')
        sppr = root.find(_q("spPr"))
        (sppr if sppr is not None else chart).addnext(tx)
    return root


def _patch_xlsx(blob):
    """xlsx 内のすべてのグラフに _fix_layout を適用した xlsx を返す。"""
    zin = zipfile.ZipFile(io.BytesIO(blob))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if re.match(r"xl/charts/chart\d+\.xml$", item.filename):
                data = etree.tostring(_fix_layout(etree.fromstring(data)), xml_declaration=True,
                                      encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
    return out.getvalue()


def _with_cache(chart_xml, ws, embed_rid):
    """グラフ XML に数値キャッシュと埋め込みブックへの参照を加える（Word が描画に使う）。"""
    root = _fix_layout(etree.fromstring(chart_xml))
    for ref in root.iter(_q("numRef")):
        if ref.find(_q("numCache")) is not None:
            continue
        vals = _ref_values(ws, ref.find(_q("f")).text)
        cache = etree.SubElement(ref, _q("numCache"))
        etree.SubElement(cache, _q("formatCode")).text = "General"
        etree.SubElement(cache, _q("ptCount")).set("val", str(len(vals)))
        for i, v in enumerate(vals):
            if isinstance(v, (int, float)):
                pt = etree.SubElement(cache, _q("pt"))
                pt.set("idx", str(i))
                etree.SubElement(pt, _q("v")).text = repr(float(v))
    ext = etree.SubElement(root, _q("externalData"))
    ext.set(f"{{{R_NS}}}id", embed_rid)
    etree.SubElement(ext, _q("autoUpdate")).set("val", "0")
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def native_chart(kind, R, D, m=None, embed_rid="rId1"):
    """Word に埋め込むための (埋め込み用xlsx, グラフXML, (幅cm, 高さcm)) を返す。
    kind: 'cool'（金属 m の冷却曲線）, 'cal'（検定）, 'dev'（規格との比較）"""
    wb = Workbook()
    wb.remove(wb.active)
    if kind == "cool":
        ws, ch = sheet_cooling(wb, R, D, m, anchor_chart=False)
        anchor = "H3"
    elif kind == "cal":
        ws, ch = sheet_lsq(wb, R, anchor_chart=False)
        anchor = "O3"
    else:
        ws, ch = sheet_dev(wb, R, anchor_chart=False)
        anchor = "I3"
    size = (ch.width, ch.height)
    ws.add_chart(ch, anchor)
    bio = io.BytesIO()
    wb.save(bio)
    xlsx = _patch_xlsx(bio.getvalue())
    chart_xml = zipfile.ZipFile(io.BytesIO(xlsx)).read("xl/charts/chart1.xml")
    return xlsx, _with_cache(chart_xml, ws, embed_rid), size
